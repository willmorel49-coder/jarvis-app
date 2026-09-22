# -*- coding: utf-8 -*-
"""Prix (PPHT) + stock disponible PAR ÉTABLISSEMENT (NR) pour le Marketing.
Lit JARVIS/PRIX ET STOCKS ETABLISSEMENTS/*_extrait.xlsx (7 établissements)
→ crm/v2/etab-prices-data.js (window.ETAB_PRICES).
Sert aux fiches marketing par établissement (bon prix + stock réel) et complète
les listes best-sellers. 100% local, Python 3.9.
Complété par la dernière extraction SOP « STOCK *.xls » du même dossier (remboursables
compris, codes CIP7, 5 sites) : pour ces sites, elle remplace les *_extrait.xlsx produit
par produit ; les produits absents gardent leur valeur. Lecture .xls : xlrd
(/usr/bin/python3).
Puis l'extraction d'un site seul « stock <SITE> <jjmmaaaa>.xlsx » (SEP le 17/09/2026 :
codes-barres CIP13, stockdispo, afmcode ; POS le 21/09/2026), même règle, la plus récente par site.
UN PRODUIT = UN CODE (21/09/2026) : le relevé des 5 sites range par CIP, POS/SEP et le catalogue par
code-barres, souvent celui du FABRICANT (3700…, 4028…) — le même produit était coupé en deux
(Versol 1000 mL : 5 315 + 1 853). Pont sûr = une ligne d'article qui porte les deux codes ;
voir canon() plus bas.
"""
import openpyxl, glob, os, json, re

SRC = '/Users/williammorel/JARVIS/PRIX ET STOCKS ETABLISSEMENTS'
OUT = 'crm/v2/etab-prices-data.js'

# code établissement = préfixe du nom de fichier (CPR, HP, MSP, OPS, POS, SEP, SOP)
def etab_code(fn):
    base = os.path.basename(fn)
    m = re.match(r'([A-Za-z]{2,4})', base)
    return m.group(1).upper() if m else base[:3].upper()

def cip_of(v):
    if v is None:
        return ''
    try:
        return str(int(float(v)))
    except (TypeError, ValueError):
        s = re.sub(r'[^0-9]', '', str(v))
        return s

def num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return 0.0

def cip13_of(c7):
    b = '34009' + c7
    s = sum(int(d) * (1 if i % 2 == 0 else 3) for i, d in enumerate(b))
    return b + str((10 - s % 10) % 10)

CAT_JS = 'crm/v2/catalogue-complet-data.js'
connus = set()  # codes du catalogue (rangé par code-barres, comme les ventes)
if os.path.exists(CAT_JS):
    connus = set(re.findall(r'\["(\d+)",', open(CAT_JS, encoding='utf-8').read()))

# Pont code-barres <-> CIP, lu dans les fichiers qui portent LES DEUX colonnes sur la même ligne
# (extractions par site + fichier catalogue « stock et prix »), jamais déduit des noms.
# Code retenu = le code-barres (catalogue et ventes sont rangés dessus : mesuré, 2 exceptions
# sur 1 312), sauf si le catalogue ne connaît que le CIP. Un code-barres porté par deux articles,
# ou un article à deux codes-barres, est AMBIGU (72 cas, dont « test d'ovulation » et « test de
# grossesse » sous le même code-barres) : on n'y touche pas, liste à arbitrer à la main.
alias = {}
def lire_pont():
    paires = set()
    srcs = glob.glob(os.path.join(SRC, 'stock *.xlsx')) + glob.glob('STATS/stock et prix*.xlsx')
    for fn in sorted(srcs):
        wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
        it = wb.active.iter_rows(values_only=True)
        ix = {str(h or '').strip().lower(): i for i, h in enumerate(next(it))}
        if 'artcode' not in ix or 'artcodebarre' not in ix:
            continue
        for r in it:
            raw = str(r[ix['artcode']] or '').strip()
            bc = cip_of(r[ix['artcodebarre']])
            if len(raw) == 7 and raw.isdigit() and len(bc) >= 8 and cip13_of(raw) != bc:
                paires.add((bc, cip13_of(raw)))
        wb.close()
    nb, nc = {}, {}
    for bc, c in paires:
        nb[bc] = nb.get(bc, 0) + 1
        nc[c] = nc.get(c, 0) + 1
    for bc, c in paires:
        if nb[bc] > 1 or nc[c] > 1:
            continue
        if c in connus and bc not in connus:
            alias[bc] = c
        else:
            alias[c] = bc
    print('  pont code-barres/CIP : %d paires, %d rapprochées, %d ambiguës laissées telles quelles'
          % (len(paires), len(alias), len(paires) - len(alias)))

def canon(code):
    return alias.get(code, code)

# Le pont est ÉCRIT sur le disque (hors dépôt : STATS/ est ignoré) pour que les générateurs
# qui lisent les VENTES par site (generate_mkt_nr.py, generate_establishments.py) rangent
# chaque vente sous le même code que le stock. Sans lui, Vismed 0,18 % unidoses (66 boîtes
# vendues à HP sous le CIP 3400974936034) restait séparé de son stock (3401074936030).
PONT_JSON = 'STATS/pont-codes.json'
def ecrire_pont():
    with open(PONT_JSON, 'w', encoding='utf-8') as fh:
        json.dump({'source': 'generate_etab_prices.py — lire_pont()', 'n': len(alias), 'alias': alias},
                  fh, ensure_ascii=False, separators=(',', ':'))
    print('  pont écrit :', PONT_JSON)

prices = {}     # code -> { cip: [ppht, stock] }
labels = {}     # cip -> désignation (pour affichage éventuel)
nr_info = {}    # cip -> [désignation, laboratoire] : NR selon AFMCODE (≠ REMBSS)
tarif = {}      # cip -> tarif (atfprix) de la dernière extraction STOCK *.xls
tarif_date = ''
etabs = []

lire_pont()
ecrire_pont()
site_files = {}  # site -> extractions « stock <SITE> <jjmmaaaa>.xlsx »
for fn in sorted(glob.glob(os.path.join(SRC, '*.xlsx'))):
    m = re.match(r'stock ([A-Za-z]{2,4}) (\d{2})(\d{2})(\d{4})\.xlsx$', os.path.basename(fn), re.I)
    if m:
        site_files.setdefault(m.group(1).upper(), []).append((m.group(4) + m.group(3) + m.group(2), fn))
        continue
    code = etab_code(fn)
    wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    hdr = [str(c) if c is not None else '' for c in next(it)]
    ix = {h: i for i, h in enumerate(hdr)}
    ci = ix.get('ARTCODEBARRE'); pi = ix.get('PPHT'); si = ix.get('STOCKDISPO')
    di = ix.get('ARTDESIGNATION'); li = ix.get('ARTCOLLECTION'); ai = ix.get('AFMCODE')
    d = prices.setdefault(code, {})
    n = 0
    for r in it:
        cip = canon(cip_of(r[ci])) if ci is not None else ''
        if not cip or len(cip) < 8:
            continue
        ppht = num(r[pi]) if pi is not None else 0.0
        stock = 0
        if si is not None:
            try: stock = int(float(r[si] or 0))
            except (TypeError, ValueError): stock = 0
        # on garde même si ppht=0 (stock utile) mais on ignore les lignes vides
        if ppht <= 0 and stock == 0:
            continue
        d[cip] = [ppht, stock]
        if di is not None and cip not in labels and r[di]:
            labels[cip] = str(r[di]).strip()
        if ai is not None and r[ai] and str(r[ai]).strip() != 'REMBSS' and cip.isdigit() and len(cip) <= 14:
            nr_info.setdefault(cip, [labels.get(cip, ''), str(r[li]).strip() if li is not None and r[li] else ''])
        n += 1
    wb.close()
    etabs.append({'code': code, 'n': n})
    print('  %s : %d produits' % (code, n))

# extraction SOP la plus récente (STOCK <jjmmaa>.xls) : artcode, artdesignation, atfprix,
# stocklivrablesop, stockmsp, stockhp, stockcpr, stockops
def date_of(fn):
    m = re.search(r'(\d{2})(\d{2})(\d{2})', os.path.basename(fn))
    return (m.group(3), m.group(2), m.group(1)) if m else ('', '', '')

stock_files = sorted(glob.glob(os.path.join(SRC, 'STOCK*.xls')), key=date_of)
if stock_files:
    import xlrd
    fn = stock_files[-1]
    ws = xlrd.open_workbook(fn).sheet_by_index(0)
    hdr = [str(c).strip().lower() for c in ws.row_values(0)]
    ix = {h: i for i, h in enumerate(hdr)}
    cols = {'SOP': 'stocklivrablesop', 'MSP': 'stockmsp', 'HP': 'stockhp',
            'CPR': 'stockcpr', 'OPS': 'stockops'}
    n = 0
    for k in range(1, ws.nrows):
        r = ws.row_values(k)
        raw = str(r[ix['artcode']]).strip()
        # CIP7 gardé en texte : cip_of() perdrait le zéro de tête (0578005)
        code = cip13_of(raw) if len(raw) == 7 and raw.isdigit() else cip_of(r[ix['artcode']])
        if len(code) < 8:
            continue
        code = canon(code)
        ppht = num(r[ix['atfprix']])
        if ppht > 0:
            tarif[code] = ppht
        for etab, col in cols.items():
            try: stock = int(float(r[ix[col]] or 0))
            except (TypeError, ValueError): stock = 0
            d = prices.setdefault(etab, {})
            old = d.get(code)
            p = ppht if ppht > 0 else (old[0] if old else 0.0)
            if p <= 0 and stock == 0:
                continue
            d[code] = [p, stock]
        if code not in labels and r[ix['artdesignation']]:
            labels[code] = str(r[ix['artdesignation']]).strip()
        n += 1
    for e in etabs:
        e['n'] = len(prices.get(e['code'], {}))
    a, m_, j = date_of(fn)
    tarif_date = '20%s-%s-%s' % (a, m_, j) if a else ''
    print('  + %s : %d produits (5 sites)' % (os.path.basename(fn), n))

site_dates = {}  # site -> date de sa dernière extraction (écrans : « au jj/mm »)
if tarif_date:
    for etab in ('SOP', 'MSP', 'HP', 'CPR', 'OPS'):
        site_dates[etab] = tarif_date
for etab, lst in sorted(site_files.items()):
    ymd, fn = max(lst)
    wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
    it = wb.active.iter_rows(values_only=True)
    ix = {str(h or '').strip().lower(): i for i, h in enumerate(next(it))}
    # POS (21/09/2026) : prix en « ppht » et laboratoire en « artcollection »
    pcol = ix['atfprix'] if 'atfprix' in ix else ix['ppht']
    lcol = ix['artmarque'] if 'artmarque' in ix else ix['artcollection']
    lu = {}  # un code-barres porté par deux articles : stocks additionnés
    for r in it:
        code = cip_of(r[ix['artcodebarre']])
        if len(code) < 8:
            raw = str(r[ix['artcode']] or '').strip()
            code = cip13_of(raw) if len(raw) == 7 and raw.isdigit() else ''
        if not code:
            continue
        code = canon(code)
        try: stock = max(0, int(float(r[ix['stockdispo']] or 0)))
        except (TypeError, ValueError): stock = 0
        ppht = num(r[pcol])
        e = lu.setdefault(code, [0.0, 0])
        if ppht > 0 and e[0] <= 0:
            e[0] = ppht
        e[1] += stock
        lib = str(r[ix['artdesignation']] or '').strip()
        if lib and code not in labels:
            labels[code] = lib
        afm = str(r[ix['afmcode']] or '').strip()
        if afm and afm != 'REMBSS' and code.isdigit() and len(code) <= 14 and lib:
            labo = str(r[lcol] or '').strip()
            nr_info.setdefault(code, [labels.get(code, lib), '' if labo == '#N/A' else labo])
    wb.close()
    d = prices.setdefault(etab, {})
    for code, (p, stock) in lu.items():
        old = d.get(code)
        p = p if p > 0 else (old[0] if old else 0.0)
        if p <= 0 and stock == 0:
            continue
        d[code] = [p, stock]
    for e in etabs:
        if e['code'] == etab:
            e['n'] = len(d)
    site_dates[etab] = '%s-%s-%s' % (ymd[:4], ymd[4:6], ymd[6:])
    print('  + %s : %d produits (%s)' % (os.path.basename(fn), len(lu), etab))

# NR tenus par un établissement mais absents du catalogue (arrêté de juin) : sans eux,
# l'écran Produits en ignorait 3 225. Nature sûre : AFMCODE des extractions par site
# (le fichier STOCK *.xls, lui, ne dit pas si un produit est remboursable).
# Codes que les extractions par site ne qualifient pas : nature tirée des fichiers de ventes
# (generate_nature_afm.py), seulement pour les produits qu'un site tient.
NATURE_JSON = os.path.join(SRC, 'nature-afm.json')
if os.path.exists(NATURE_JSON):
    tenus = set(tarif).union(*[set(d) for d in prices.values()])
    for code, (afm, labo) in json.load(open(NATURE_JSON, encoding='utf-8')).items():
        code = canon(code)
        if afm != 'REMBSS' and code in tenus and code not in nr_info and labels.get(code):
            nr_info[code] = [labels[code], labo]
extra = {}
if connus:
    extra = {c: v for c, v in nr_info.items() if c not in connus and v[0]}
    print('  NR hors catalogue : %d (catalogue : %d réf.)' % (len(extra), len(connus)))

# combiné "TOUS" : meilleur PPHT (>0) + stock total sur tous les établissements
allc = {}
for code, d in prices.items():
    for cip, (p, s) in d.items():
        e = allc.setdefault(cip, [0.0, 0])
        if p > 0 and (e[0] == 0 or p < e[0]):
            e[0] = p
        e[1] += s

# Quels sites ont VENDU le produit (Appro : un transfert vers un site qui ne le vend pas se signale).
# Source : STATS/<SITE>_Pharma_agregation.xlsx (cumul des ventes du site, sans mois). On ne sort
# qu'un oui/non par site — ni quantité ni chiffre d'affaires. Le CIP est dans ARTCODE (CIP7 ou 13),
# ARTCODEBARRE étant souvent un EAN fabricant : on retient celui des deux que le stock connaît.
vendu = {}        # cip -> masque de bits : bit i = une vente est relevée sur vendu_sites[i]
vendu_sites = []  # sites dont le fichier de ventes a été lu : hors de cette liste = on ne sait pas
for site in sorted(prices):
    fn = os.path.join('STATS', '%s_Pharma_agregation.xlsx' % site)
    if not os.path.exists(fn):
        continue
    ws = openpyxl.load_workbook(fn, read_only=True, data_only=True).active
    it = ws.iter_rows(values_only=True)
    ix = {str(h): i for i, h in enumerate(next(it)) if h is not None}
    if 'QTE_TOTALE' not in ix or 'ARTCODE' not in ix:
        continue
    vendu_sites.append(site)
    for r in it:
        try:
            if float(r[ix['QTE_TOTALE']] or 0) <= 0:
                continue
        except (TypeError, ValueError):
            continue
        raw = cip_of(r[ix['ARTCODE']])
        cands = [cip13_of(raw) if len(raw) == 7 else raw]
        if 'ARTCODEBARRE' in ix:
            cands.append(cip_of(r[ix['ARTCODEBARRE']]))
        for c in cands:
            c = canon(c)
            if c in allc:
                vendu.setdefault(c, set()).add(site)
                break
vendu = {c: sum(1 << vendu_sites.index(x) for x in v) for c, v in vendu.items()}
print('  ventes par site : %d produits, sites lus : %s' % (len(vendu), ', '.join(vendu_sites) or 'aucun'))

data = {'etabs': sorted(etabs, key=lambda x: x['code']), 'prices': prices, 'all': allc,
        'tarif': tarif, 'tarifDate': tarif_date, 'siteDates': site_dates, 'nrHorsCat': extra,
        'vendu': vendu, 'venduSites': vendu_sites}
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, 'w', encoding='utf-8') as fh:
    fh.write('// Prix PPHT + stock par établissement (NR) — generate_etab_prices.py\n')
    fh.write('window.ETAB_PRICES = ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n')

print('OK ->', OUT)
print('  établissements :', ', '.join('%s(%d)' % (e['code'], e['n']) for e in data['etabs']))
print('  CIP combinés   :', len(allc))
print('  taille fichier : %.0f Ko' % (os.path.getsize(OUT) / 1024))
