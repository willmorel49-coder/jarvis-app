#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Marge « Notre marge » par client x mois x famille — fichier PROTEGE `marge-clients.js` (14/10/2026).

CE QUE FAIT CE SCRIPT
Relit les tableurs de ventes (STATS/<CODE>_MM_AAAA.xlsx + STATS/hors-perimetre/) et estime, pour
chaque officine et chaque mois, la marge que l'entreprise garde, rangee par famille d'affichage
(p, p1, p2, i, c, fr, nr). Une seule estimation : celle de la direction.

LE PRINCIPE (une cascade, du plus sur au moins sur)
La direction compte la marge ligne par ligne (prix net moins son prix d'achat). Son prix d'achat n'est
lisible que sur une partie des lignes (colonne MARGETHEORIQUE) ; il est appris sur ces lignes puis
applique aux autres :
  1. lue            : la marge est dans la ligne ;
  2. prix d'achat   : appris sur le MEME produit, le MEME mois, au MEME prix brut ;
  3. prix d'achat   : appris sur le meme produit au meme prix brut, tous mois ;
  4. taux du produit: marge / prix net appris sur le produit ;
  5. repli          : generique et non rembourse = taux appris ; petit prix d'un laboratoire a
                      contrat = montant fixe ; sinon prix net moins prix fabricant retrouve depuis
                      le prix brut (bareme des grossistes, loi) ; sans prix brut : regle du Pilotage.
Les niveaux 2 a 4 sont appris PAR REGIME (2026 / 2027) : un prix d'achat de 2026 n'est jamais applique
a une vente de 2027 (reforme du 20/01/2027, arrete du 3 aout 2026).

MOIS DE REMPLACEMENT (decision de Will, 15/10/2026)
Quand le prix d'un produit change en cours de mois, la direction a vendu au nouveau prix des boites achetees a
l'ancien : la marge du mois s'effondre sans rien dire du travail du commercial. Le reglage prive peut alors dire
« pour ces produits, ce mois-la, compter la marge par boite d'un autre mois » (juin 2026 -> juillet 2026 pour
trois produits). Marge par boite = celle que la cascade donne a l'autre mois, au meme prix brut, sinon au prix
brut le plus proche (meme code produit d'abord, puis tous les codes du produit). La mention affichee sous la
carte est fabriquee a partir de ce meme reglage : elle ne peut pas dire autre chose que le calcul.

BAISSES DE PRIX (decision de Will, 15/10/2026 : « ce n'est pas de notre fait »)
Une baisse de prix ne compte jamais dans la marge d'un commercial, et plus aucune liste n'est a taper : la baisse est
reperee toute seule, code par code, dans le meme regime. Un mois est touche quand le code s'y vend a un prix brut plus
bas d'au moins 1 % que le brut le plus vendu du mois d'avant ET que la marge par boite lue a ce brut tombe sous celle du
mois de reference (meme brut, sinon le plus proche ramene au meme taux de marge ; ecart au-dela de 0,50 EUR ou 2 %).
Les boites du code vendues ce mois-la au prix baisse sont alors comptees avec la marge par boite du mois de reference =
le premier mois suivant sans baisse (meme calcul que les mois de remplacement ; autre brut = meme taux de marge) ; celles
vendues encore a l'ancien prix gardent leur marge, et aucune boite ne descend sous sa marge ordinaire. Baisse du dernier mois lu : marge d'avant la baisse, en taux du prix net, en
attendant le mois suivant (le calcul se refait seul a chaque passage). Un code deja dans un mois de remplacement du
reglage garde ce remplacement. La mention du mois est fabriquee depuis les baisses reperees.

CE QUI NE VIT PAS DANS LE DEPOT
Rien de prive ici : ni prix d'achat, ni taux appris, ni laboratoire, ni montant de la regle privee.
Les parametres prives sont dans ~/.config/jarvis/marge-direction-prive.json (illisible = code 3, rien
n'est ecrit). Le bareme des grossistes est une loi : il reste dans le code.

SORTIE (gitignoree, JAMAIS commitee) : marge-clients.js a la racine du depot.
  window.MARGE_CLIENTS = {v:1, fait, jusqua, mois, reserves, fam, c:[prenoms], o:[pharmacyId],
                          livraison:{cout, moy:{n, jm, parJour, renta}},
                          pilotageMoy:{f:{'AAAA-MM|AAAA-MM':{n, jm, parJour, renta, caCl, mgCl, nSect, secteurs}, ...},
                                       couv:{'AAAA-MM':[secteurs complets, secteurs attendus], ...}},
                          r:[[ci, oi, ym, ca, tot, p, p1, p2, i, c, fr, nr, caGen, jours], ...]}
  ca = CA net HT de toutes les lignes (generiques compris) ; caGen = la part generique de ce CA
  (marge 0) ; tot = somme exacte des 7 familles apres arrondi au centime.
  jours = jours de commande DISTINCTS de l'officine ce mois-la, tous commerciaux confondus (meme valeur sur
  chaque ligne de l'officine-mois : la page prend le max, jamais la somme). Un jour de commande = une date
  avec au moins une facture FC (pas un avoir AC), comme la page de reunion du 28/10/2026.
  livraison (28/10/2026) = cout d'une livraison (reglage prive `cout_livraison`) + l'officine MOYENNE de TOUS
  les clients, calculee ici et recopiee telle quelle dans chaque jeu : un commercial ne charge que ses
  officines, une moyenne faite dans la page serait celle de SES clients. Cout absent ou illisible = pas de cle.
  pilotageMoy (28/10/2026) = la meme officine moyenne pour CHAQUE fenetre [du, au] de mois du fichier (45 fenetres sur 9 mois),
  pour le bloc « Mon secteur compare a la moyenne Integral » du Pilotage : calculee ici sur le jeu COMPLET, recopiee telle
  quelle dans chaque jeu, jamais recalculee dans la page. Presente seulement si livraison l'est (la rentabilite en depend).
decouper_par_commercial.py en tire un fichier par jeu de commercial.

Usage : /usr/bin/python3 generate_marge_clients.py
Codes de sortie : 0 ok · 2 entree introuvable · 3 reglage prive illisible · 4 rien a apprendre ·
5 ecriture non verifiee.
"""
import ast
import datetime
import glob
import gzip
import hashlib
import json
import math
import os
import re
import subprocess
import sys
import time
import unicodedata
import warnings
from collections import defaultdict
from multiprocessing import Pool

warnings.filterwarnings('ignore')
sys.dont_write_bytecode = True

BASE = os.path.dirname(os.path.abspath(__file__))
SORTIE = os.path.join(BASE, 'marge-clients.js')
PRIVE = os.environ.get('JARVIS_MARGE_DIRECTION_PRIVE') or os.path.expanduser('~/.config/jarvis/marge-direction-prive.json')
CACHE = os.path.expanduser('~/.cache/jarvis-marge-clients')
ENTETE = '// ⚠️ NE JAMAIS COMMITER. Servi par adresse signée (Supabase).'
VERSION_CACHE = 'g3'                 # g3 (28/10/2026) : le cache garde aussi les jours de commande

# ── Barème des grossistes (la LOI, publique) ───────────────────────────────────────────────────
# Prix grossiste (« brut », PGHT) = prix fabricant (PFHT) + max(plancher, 6,93 % du PFHT), la part
# proportionnelle étant plafonnée à 32,4996 € (6,93 % de 468,97 €). Le forfait froid s'ajoute au brut.
# Réforme du 20/01/2027 (arrêté du 03/08/2026) : plancher 0,30 -> 0,22 €, forfait froid 0,63 -> 0,80 €.
ANOM_JOUR = ('2269223', '2026-04-30')  # jour anormal écarté, comme la page de réunion du 28/10/2026
TAUX_GROSSISTE = 0.0693
PFHT_PLAFOND = 468.97
PLAFOND_PROPORTIONNEL = 32.4996                        # = 6,93 % de 468,97 €, arrondi comme le bareme
BRUT_PLAFOND = PFHT_PLAFOND + PLAFOND_PROPORTIONNEL    # 501,4696 : au-dela, brut - 32,4996
BAREME = {
    2026: {'plancher': 0.30, 'froid': 0.63},
    2027: {'plancher': 0.22, 'froid': 0.80},
}
for _r in BAREME.values():
    # Le plancher l'emporte tant que 6,93 % du PFHT lui reste inférieur : bascule du PFHT = plancher / 6,93 %
    # (2026 : 4,329 €, 2027 : 3,1746 €) ; en brut, la bascule est « bascule + plancher » (2026 : 4,629 €,
    # 2027 : 3,3946 €). Calculé, jamais recopié : changer le plancher recalcule les deux bornes.
    _r['bascule_pfht'] = _r['plancher'] / TAUX_GROSSISTE
    _r['seuil_brut'] = _r['bascule_pfht'] + _r['plancher']

# ── Règle publique du Pilotage (v2-pilotage.js, MARGE) : sert quand on ne connait pas le prix brut ──
PETIT, TAUX_INTER, CHER, COEF_NR = 0.12, 0.0304, 13.0, 1.03
PALIER_PETIT, PALIER_CHER = 4.33, 468.0
# 4e tranche de 2027 (Pilotage, marge_2027) : < 3,1746 € -> 0,088 €/boîte ; 3,1746 à 4,33 € -> 2,772 % du net.
BASCULE_2027, PETIT_2027, TAUX_ENTRE_2027 = 3.1746, 0.088, 0.02772
DEBUT_2027 = (2027, 1, 20)    # jour d'entrée en vigueur : une ligne datée à partir de là est du régime 2027
MOIS_2027 = 202702            # sans date de ligne : tout mois à partir de février 2027

FAM = ['p', 'p1', 'p2', 'i', 'c', 'fr', 'nr']
MOIS_FR = ['janvier', 'février', 'mars', 'avril', 'mai', 'juin', 'juillet', 'août', 'septembre', 'octobre', 'novembre', 'décembre']
SOURCES_ENV = ('SOURCES', 'RATTACHEES', 'AVOIRS_EN_POSITIF', 'MOIS_LUS')


# ══════════════════════════════════════════════ fonctions pures (testées) ════════════════════════
def sa(v):
    """Minuscules sans accents, pour comparer des libellés."""
    return ''.join(c for c in unicodedata.normalize('NFD', str(v or '').strip().lower()) if unicodedata.category(c) != 'Mn')


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def cip13(v):
    if v is None:
        return ''
    s = str(v).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return re.sub(r'\D', '', s)


def regime_ligne(date, ym):
    """2027 si la ligne est datée (PCVDATEEFFET) à partir du 20/01/2027 ; une ligne datée avant reste 2026 ;
    sans date lisible, on regarde le mois du fichier (février 2027 et après)."""
    if hasattr(date, 'year'):
        t = (date.year, date.month, date.day)
    else:
        m = re.match(r'\s*(\d{4})-(\d{2})-(\d{2})', str(date or ''))
        t = tuple(int(x) for x in m.groups()) if m else None
    if t:
        return 2027 if t >= DEBUT_2027 else 2026
    return 2027 if ym >= MOIS_2027 else 2026


def pf(B, froid, regime):
    """Prix fabricant retrouvé depuis le prix brut (barème des grossistes du régime), forfait froid retiré."""
    p = BAREME[regime]
    b = B - (p['froid'] if froid else 0.0)
    if b <= p['seuil_brut']:
        return b - p['plancher']
    if b <= BRUT_PLAFOND:
        return b / (1 + TAUX_GROSSISTE)
    return b - PLAFOND_PROPORTIONNEL


def ym_de(t):
    """'2026-06' -> 202606. Lève AttributeError si le texte n'a pas cette forme."""
    a, m = re.fullmatch(r'(\d{4})-(0[1-9]|1[0-2])', str(t)).groups()
    return int(a) * 100 + int(m)


def mention(noms, ym, ym_par, motif=''):
    """Texte affiché après « Juin : » (Pilotage) et « Juin 2026 : » (fiche officine), tiré du réglage."""
    liste = noms[0] if len(noms) == 1 else ', '.join(noms[:-1]) + ' et ' + noms[-1]
    mois = MOIS_FR[ym_par % 100 - 1] + ('' if ym_par // 100 == ym // 100 else ' %d' % (ym_par // 100))
    t = '%s %s avec %s marge par boîte de %s' % (liste, 'compté' if len(noms) == 1 else 'comptés', 'sa' if len(noms) == 1 else 'leur', mois)
    return t + (' (%s)' % motif if motif else '')


def lire_remplacements(liste):
    """Bloc `remplacements` du réglage privé -> {ym: {'par': ym, 'cips': {cip: produit}, 'noms': [...], 'texte': mention}}.
    Lève une exception si le bloc est mal formé : mieux vaut ne rien écrire que compter un mois de travers."""
    out = {}
    for x in liste or []:
        ym, par = ym_de(x['mois']), ym_de(x['par'])
        noms = [str(n) for n in x['produits']]
        cips = {}
        for n in noms:
            codes = [cip13(c) for c in x['produits'][n]]
            assert codes and all(len(c) == 13 for c in codes), 'codes produit de %s' % n
            for c in codes:
                assert c not in cips, 'code %s dans deux produits' % c
                cips[c] = n
        assert noms and par != ym and ym not in out
        out[ym] = {'par': par, 'cips': cips, 'noms': noms, 'texte': mention(noms, ym, par, str(x.get('motif') or ''))}
    assert not set(out) & {r['par'] for r in out.values()}, 'un mois remplacé sert aussi de remplacement'
    return out


def par_boite(tables, cle, nom, B):
    """Marge par boîte du mois de remplacement pour un groupe de lignes (code `cle`, produit `nom`, prix brut `B`).
    tables = ({code: {brut: [marge, boîtes]}}, {produit: {brut: [marge, boîtes]}}). Ordre : même code au même brut,
    même produit au même brut, même code au brut le plus proche, même produit au brut le plus proche.
    Rend (marge par boîte, niveau), ou (None, None) si le produit n'a pas été vendu ce mois-là."""
    a, b = tables[0].get(cle) or {}, tables[1].get(nom) or {}
    for t, niveau in ((a, 'code'), (b, 'produit')):
        if B in t:
            return t[B][0] / t[B][1], niveau + ' même brut'
    for t, niveau in ((a, 'code'), (b, 'produit')):
        if t:
            k = min(t, key=lambda x: (abs(x - B), x))
            return t[k][0] / t[k][1], niveau + ' brut proche'
    return None, None


SEUIL_BAISSE = 0.99              # brut au moins 1 % sous le brut le plus vendu du mois d'avant
TOL_EUROS, TOL_TAUX = 0.50, 0.02  # marge par boîte plus basse que la référence au-delà de 0,50 € ou 2 % = perte


def rang(ym):
    return (ym // 100) * 12 + ym % 100


def nom_mois(ym, ym_texte):
    """« juillet », « avril 2027 » (année écrite si elle diffère de celle du mois dont on parle)."""
    return MOIS_FR[ym % 100 - 1] + ('' if ym // 100 == ym_texte // 100 else ' %d' % (ym // 100))


def de(mot):
    return ("d'" if mot[:1] in 'aeiouéèêâô' else 'de ') + mot


def detecter_baisses(om, ventes, exclus=None):
    """Baisses de prix repérées dans les ventes.
    om = {(régime, code, ym, brut): [marge lue, boîtes, q×net]} (Cascade.OM, lignes dont la marge est lue) ;
    ventes = {(régime, code): {ym: {brut: boîtes}}} (toutes les ventes positives) ; exclus = {ym: codes d'un mois de
    remplacement du réglage}. Rend {ym: {code: (mois de référence, taux, brut d'avant)}} : mois de référence = premier mois
    suivant lu sans baisse ni remplacement, taux None ; sans lui : (None, marge / net d'avant la baisse). Brut d'avant =
    le plus vendu du mois précédent : seules les boîtes vendues au moins 1 % en dessous portent la baisse."""
    exclus = exclus or {}
    lue = defaultdict(lambda: defaultdict(dict))
    for (reg, cle, ym, B), x in om.items():
        if x[1] > 0 and x[2] > 0:
            lue[(reg, cle)][ym][B] = x
    out = defaultdict(dict)
    for (reg, cle), v in ventes.items():
        L = lue.get((reg, cle))
        if not L:
            continue
        mois = sorted(v)
        exclu = {m for m in mois if cle in exclus.get(m, ())}
        for i in range(len(mois) - 1, 0, -1):            # du plus récent au plus ancien : la référence est déjà jugée
            m, p = mois[i], mois[i - 1]
            if m in exclu or not L.get(m) or rang(m) - rang(p) > 2 or not v[p]:
                continue
            bp = max(v[p], key=lambda b: (v[p][b], b))
            n = next((k for k in mois[i + 1:] if L.get(k) and k not in exclu and cle not in out.get(k, {})), None)
            taux = None
            if n is None:
                av = L.get(p) or {}
                xs = [av[bp]] if bp in av else list(av.values())
                if not xs or not sum(x[2] for x in xs):
                    continue
                taux = sum(x[0] for x in xs) / sum(x[2] for x in xs)
            for B, x in L[m].items():
                if B >= bp * SEUIL_BAISSE:
                    continue
                if n is None:
                    r = taux * x[2] / x[1]
                else:
                    t = L[n]
                    k = B if B in t else min(t, key=lambda b: (abs(b - B), b))
                    r = t[k][0] / t[k][1] * B / k        # autre brut : même taux de marge
                if x[0] / x[1] < r - max(TOL_EUROS, TOL_TAUX * abs(r)):
                    out[m][cle] = (n, taux, bp)
                    break
    return dict(out)


def marge_baisse(g, info, tables, sans):
    """Marge imposée à un groupe de boîtes vendues au prix baissé. info = (mois de référence, taux, brut d'avant) ;
    tables = celles des mois de référence ; sans = marge ordinaire du groupe. Mois de référence connu : marge par boîte
    à ce mois (même brut, sinon le plus proche ramené au même taux de marge) ; sinon taux d'avant la baisse × net.
    Rend (None, niveau) quand la marge ordinaire est déjà au moins à ce niveau : la règle enlève une perte, rien d'autre."""
    n, taux, _bp = info
    B = g[3]
    if n is None:
        imp, niveau = g[12] * taux, "taux d'avant la baisse"
    else:
        pb, niveau = par_boite(tables[n], g[2], g[2], B)
        if pb is None:
            return None, None
        if niveau.endswith('brut proche'):
            t = tables[n][0].get(g[2]) or tables[n][1].get(g[2])
            k = min(t, key=lambda x: (abs(x - B), x))
            pb = pb * B / k
        imp = g[11] * pb
    if imp <= sans:
        return None, 'marge ordinaire déjà au niveau'
    return imp, niveau


def texte_baisses(d, ym, autres=False):
    """Mention d'un mois touché par des baisses de prix ({code: (mois de référence, taux)})."""
    nb, refs = len(d), {v[0] for v in d.values()}
    sujet = '%d %sproduit%s dont le prix a baissé compté%s' % (nb, 'autre ' if autres and nb == 1 else 'autres ' if autres else '',
                                                             '' if nb == 1 else 's', '' if nb == 1 else 's')
    leur = 'sa' if nb == 1 else 'leur'
    if refs == {None}:
        suite = ym + 1 if ym % 100 < 12 else (ym // 100 + 1) * 100 + 1
        fin = '%s marge d\'avant la baisse, en attendant les ventes %s' % (leur, de(nom_mois(suite, ym)))
    elif len(refs) == 1:
        fin = '%s marge par boîte %s' % (leur, de(nom_mois(refs.pop(), ym)))
    else:
        fin = '%s marge par boîte du mois d\'après la baisse' % leur
    return '%s avec %s : la perte sur le stock n\'est pas comptée' % (sujet, fin)


class Listes:
    """Listes de produits (CIP13) et montants. Rien d'autre n'entre dans le calcul."""

    def __init__(self, gen=(), nr=(), fr=(), privee=(), petit_su=0.0, petit_su_2027=0.0, labos=()):
        self.gen, self.nr, self.fr, self.privee = set(gen), set(nr), set(fr), set(privee)
        self.petit_su, self.petit_su_2027 = float(petit_su), float(petit_su_2027)
        self.labos = tuple(x.upper() for x in labos)

    def petit(self, regime):
        return self.petit_su_2027 if regime == 2027 else self.petit_su


def famille(cip, N, regime, L, rb=True):
    """Famille d'affichage EXCLUSIVE, par priorité : générique (None, marge 0) > nr > fr > prix net unitaire.
    « Non remboursé » = la ligne (ou le catalogue) n'est pas remboursable à TVA 2,1 % : même critère que la cascade
    de la direction (la liste Sagitta n'est plus dans le dépôt, 14/10/2026)."""
    if cip in L.gen:
        return None
    if not rb:
        return 6
    if cip in L.fr:
        return 5
    if N < PALIER_PETIT:
        if regime == 2027:
            return 1 if N < BASCULE_2027 else 2
        return 0
    return 3 if N < PALIER_CHER else 4


def ecran(g, L):
    """Repli quand le prix brut est inconnu (0) : règle publique du Pilotage, datée, + liste privée (montant fixe)."""
    off, reg, cle, B, N, rb, froid, gen, su, pos, a_mt, q, net, qN, mt = g
    cip = '' if cle.startswith('#') else cle
    if cip in L.gen:
        return 0.0
    if not rb:
        m = net * (1 - 1 / COEF_NR)
    elif N < PALIER_PETIT:
        if cip in L.privee:
            m = q * L.petit(reg)
        elif reg == 2027:
            m = q * PETIT_2027 if N < BASCULE_2027 else net * TAUX_ENTRE_2027
        else:
            m = q * PETIT
    elif N < PALIER_CHER:
        m = net * TAUX_INTER
    else:
        m = q * CHER
    if cip in L.fr:
        m += q * BAREME[reg]['froid']
    return m


class Cascade:
    """Prix d'achat de la direction, appris ligne à ligne puis appliqué aux autres lignes, par régime."""

    def __init__(self, L):
        self.L = L
        self.OM, self.OB, self.OC = {}, {}, {}   # (régime, produit, mois, brut) · (régime, produit, brut) · (régime, produit) -> [marge, boîtes, q×net]
        self.cat = {'gen': [0.0, 0.0], 'nr': [0.0, 0.0], 'bar': [0.0, 0.0]}
        self.TG = self.TNR = 0.0
        self.lignes_connues = 0

    def apprendre(self, g, ym):
        off, reg, cle, B, N, rb, froid, gen, su, pos, a_mt, q, net, qN, mt = g
        if not a_mt:
            return
        self.lignes_connues += 1
        k = 'gen' if gen else 'nr' if not rb else 'bar'
        self.cat[k][0] += mt
        self.cat[k][1] += net
        if not pos or not q:
            return
        for d, key in ((self.OM, (reg, cle, ym, B)), (self.OB, (reg, cle, B)), (self.OC, (reg, cle))):
            x = d.get(key)
            if x is None:
                x = d[key] = [0.0, 0.0, 0.0]
            x[0] += mt
            x[1] += q
            x[2] += qN

    def figer(self):
        """Taux appris des génériques et des non remboursés (repli). Faux = rien à apprendre."""
        if not self.cat['gen'][1] or not self.cat['nr'][1]:
            return False
        self.TG = self.cat['gen'][0] / self.cat['gen'][1]
        self.TNR = self.cat['nr'][0] / self.cat['nr'][1]
        return True

    @staticmethod
    def _somme(d, key):
        v = d.get(key)
        return v if v and v[1] > 0 and v[2] > 0 else None

    def marge(self, g, ym):
        """(marge, provenance) d'un groupe de lignes. Provenance : lue · prix d'achat · taux du produit · repli · ecran."""
        L = self.L
        off, reg, cle, B, N, rb, froid, gen, su, pos, a_mt, q, net, qN, mt = g
        cip = '' if cle.startswith('#') else cle
        if cip in L.gen:
            return 0.0, 'gen'
        if a_mt:
            return mt, 'lue'
        for d, key in ((self.OM, (reg, cle, ym, B)), (self.OB, (reg, cle, B))):
            s = self._somme(d, key)
            if s:
                return qN - q * (s[2] - s[0]) / s[1], "prix d'achat"
        s = self._somme(self.OC, (reg, cle))
        if s:
            return net * s[0] / s[2], 'taux du produit'
        if gen:
            return net * self.TG, 'repli'
        if not rb:
            return net * self.TNR, 'repli'
        if su and N < PALIER_PETIT:
            return q * L.petit(reg), 'repli'
        if B:
            return qN - q * pf(B, froid, reg), 'repli'
        return ecran(g, L), 'ecran'


def ajouter(lignes, cle, g, ym, cascade, L, marge=None):
    """Ajoute un groupe de lignes à l'agrégat `lignes[(prénom, officine, ym)]` = 7 familles + ca + caGen.
    `marge` imposée = mois de remplacement. Rend la provenance de la marge (pour la mesure)."""
    off, reg, cle_p, B, N, rb, froid, gen, su, pos, a_mt, q, net, qN, mt = g
    cip = '' if cle_p.startswith('#') else cle_p
    x = lignes.get(cle)
    if x is None:
        x = lignes[cle] = [0.0] * 9      # 0..6 familles · 7 = ca · 8 = caGen
    x[7] += net
    f = famille(cip, N, reg, L, rb)
    if f is None:
        x[8] += net
        return 'gen', 0.0
    m, src = cascade.marge(g, ym) if marge is None else (marge, 'mois de remplacement')
    x[f] += m
    return src, m


def arrondir(lignes, c_idx, o_idx):
    """Lignes de sortie : [ci, oi, ym, ca, tot, p, p1, p2, i, c, fr, nr, caGen], montants au centime, tot = somme des 7 familles ARRONDIES."""
    out = []
    for (prenom, officine, ym), x in lignes.items():
        fam = [round(v, 2) for v in x[:7]]
        tot = round(sum(fam), 2)
        ca, cag = round(x[7], 2), round(x[8], 2)
        if ca == 0 and tot == 0 and cag == 0 and not any(fam):
            continue
        out.append([c_idx[prenom], o_idx[officine], ym, ca, tot] + fam + [cag])
    out.sort(key=lambda r: (r[0], r[1], r[2]))
    return out


# ───────────────────────────────────────────── format du fichier ─────────────────────────────────
def formater(d):
    """Texte du fichier. Une clé par ligne : `analyser` (ici et dans decouper_par_commercial) le relit sans moteur JS."""
    j = lambda v: json.dumps(v, ensure_ascii=False, separators=(',', ':'))
    lignes = [ENTETE, "window.MARGE_CLIENTS = {v:%d, fait:'%s', jusqua:'%s'," % (d['v'], d['fait'], d['jusqua']),
              'mois:%s,' % j(d['mois']), 'reserves:%s,' % j(d['reserves']), 'fam:%s,' % j(d['fam']),
              'c:%s,' % j(d['c']), 'o:%s,' % j(d['o'])]
    if 'livraison' in d:
        lignes.append('livraison:%s,' % j(d['livraison']))
    if 'pilotageMoy' in d:
        lignes.append('pilotageMoy:%s,' % j(d['pilotageMoy']))
    lignes.append('r:[')
    lignes.append(',\n'.join(j(r) for r in d['r']))
    lignes.append(']};')
    return '\n'.join(lignes) + '\n'


def analyser(texte):
    """Inverse de `formater`. Lève ValueError si le fichier n'a pas la forme attendue."""
    L = texte.split('\n')
    try:
        i = next(k for k, l in enumerate(L) if l.startswith('window.MARGE_CLIENTS = {v:'))
        m = re.match(r"window\.MARGE_CLIENTS = \{v:(\d+), fait:'([^']*)', jusqua:'([^']*)',$", L[i])
        d = {'v': int(m.group(1)), 'fait': m.group(2), 'jusqua': m.group(3)}
        for k, cle in enumerate(('mois', 'reserves', 'fam', 'c', 'o'), 1):
            assert L[i + k].startswith(cle + ':') and L[i + k].endswith(',')
            d[cle] = json.loads(L[i + k][len(cle) + 1:-1])
        k = i + 6
        if L[k].startswith('livraison:'):        # facultative : un fichier d'avant le 28/10/2026 n'en a pas
            assert L[k].endswith(',')
            d['livraison'] = json.loads(L[k][len('livraison:'):-1])
            k += 1
        if L[k].startswith('pilotageMoy:'):      # facultative aussi (28/10/2026, bloc « Mon secteur » du Pilotage)
            assert L[k].endswith(',')
            d['pilotageMoy'] = json.loads(L[k][len('pilotageMoy:'):-1])
            k += 1
        assert L[k] == 'r:['
        fin = next(n for n in range(len(L) - 1, i, -1) if L[n] == ']};')
        corps = L[k + 1:fin]
        d['r'] = [json.loads(l.rstrip(',')) for l in corps if l.strip()]
    except (StopIteration, AssertionError, AttributeError, IndexError, json.JSONDecodeError) as e:
        raise ValueError('marge-clients.js illisible (%s)' % e)
    return d


# ═══════════════════════════════════════════ entrées (disque) ════════════════════════════════════
def lire_prive():
    """Réglage privé, ou None s'il est illisible."""
    try:
        r = json.load(open(PRIVE, encoding='utf-8'))
        labos = [str(x) for x in r['laboratoires_contenant']]
        petit = float(r['petit_par_boite'])
        petit27 = float(r.get('petit_par_boite_2027', petit))
        res = {str(k): str(v) for k, v in (r.get('reserves') or {}).items()}
        rempl = lire_remplacements(r.get('remplacements'))
        for ym, x in rempl.items():      # la mention d'un mois remplacé vient du remplacement lui-même
            res['%d-%02d' % (ym // 100, ym % 100)] = x['texte']
        assert labos and petit >= 0 and petit27 >= 0
        cout = r.get('cout_livraison')       # facultatif : absent ou illisible = pas de bloc « Rentabilité », jamais une valeur en dur
        cout = float(cout) if isinstance(cout, (int, float)) and not isinstance(cout, bool) and cout > 0 else None
        return {'labos': labos, 'petit': petit, 'petit_2027': petit27, 'reserves': res, 'remplacements': rempl, 'cout_livraison': cout}
    except Exception:
        return None


def jours_par_mois(couples):
    """Couples (officine, 'AAAA-MM-JJ') -> {(officine, AAAAMM): nombre de dates distinctes}. Une officine servie par deux
    commerciaux le même jour compte UN jour. Écarte le compte INTER et le jour anormal ANOM_JOUR."""
    J = defaultdict(int)
    for off, d in set(map(tuple, couples)):
        if off == 'INTER' or (off, d) == ANOM_JOUR or not re.match(r'^\d{4}-\d{2}-\d{2}$', d):
            continue
        J[(off, int(d[:4] + d[5:7]))] += 1
    return dict(J)


def arrondi_js(x):
    """Math.round(x * 100) / 100 de la page (round() de Python arrondit au pair)."""
    return math.floor(x * 100 + 0.5) / 100


def moyenne_livraison(d, jours, cout):
    """L'officine moyenne de TOUS les clients de `d` (même calcul que la fiche, mesuré dans Safari le 28/10/2026) :
    officines qui ont des lignes ET des jours de commande sur les mois de d['mois'] ; par officine, jm = jours / nm,
    parJour = marge / jours, renta = (marge - jours x coût) / nm, nm = mois avec au moins un jour ; puis moyenne simple.
    `jours` = {(officine, AAAAMM): n}. Rend None s'il n'y a aucune officine."""
    mois = [int(m[:4] + m[5:7]) for m in d['mois']]
    dans = set(mois)
    tot = {}
    for r in d['r']:
        if r[2] in dans:
            o = str(d['o'][r[1]])
            tot[o] = tot.get(o, 0.0) + r[4]
    L = []
    for o, t in tot.items():
        t = arrondi_js(t)
        js = [jours.get((o, m), 0) for m in mois]
        n, nm = sum(js), sum(1 for x in js if x > 0)
        if n:
            L.append((n / nm, t / n, (t - n * cout) / nm))
    if not L:
        return None
    return {'n': len(L), 'jm': round(sum(x[0] for x in L) / len(L), 5), 'parJour': round(sum(x[1] for x in L) / len(L), 5),
            'renta': round(sum(x[2] for x in L) / len(L), 5)}


def moyennes_pilotage(d, cout):
    """28/10/2026 — l'officine moyenne du RÉSEAU ENTIER pour chaque fenêtre [du, au] de mois de d['mois'] (45 sur 9 mois).
    Même définition que moyenne_livraison, par fenêtre : par officine, sur les mois où elle a commandé (jours > 0),
    jm = jours / nm, parJour = marge / jours, renta = (marge - jours x coût) / nm, caCl = CA / nm, mgCl = marge / nm ;
    moyenne simple des officines. `jours` d'un officine-mois = le max des lignes (même valeur sur chacune).
    nSect = clients actifs d'un « secteur moyen » (officines actives comptées dans chaque secteur où elles figurent).
    Rend {'f': {'AAAA-MM|AAAA-MM': {...} ou None}} ; la fenêtre entière doit retomber sur livraison.moy (contrôlé à l'appelant)."""
    mois = d['mois']
    K = [int(m[:4] + m[5:7]) for m in mois]
    fen = {}
    for a in range(len(K)):
        for b in range(a, len(K)):
            off = {}
            for r in d['r']:
                if not K[a] <= r[2] <= K[b]:
                    continue
                x = off.setdefault(str(d['o'][r[1]]), {'ca': 0.0, 'mg': 0.0, 'j': {}, 'ci': set()})
                x['ca'] += r[3]
                x['mg'] += r[4]
                x['ci'].add(r[0])
                if (r[13] or 0) > x['j'].get(r[2], 0):
                    x['j'][r[2]] = r[13]
            P = []
            par_sect = {}
            for x in off.values():
                n = sum(x['j'].values())
                nm = len([v for v in x['j'].values() if v > 0])
                if n:
                    P.append((n / nm, x['mg'] / n, (x['mg'] - n * cout) / nm, x['ca'] / nm, x['mg'] / nm))
                    for ci in x['ci']:
                        par_sect[ci] = par_sect.get(ci, 0) + 1
            if not P:
                fen[mois[a] + '|' + mois[b]] = None
                continue
            k = len(P)
            moy = lambda q: round(sum(p[q] for p in P) / k, 4)
            fen[mois[a] + '|' + mois[b]] = {'n': k, 'jm': moy(0), 'parJour': moy(1), 'renta': moy(2), 'caCl': moy(3), 'mgCl': moy(4),
                                            'nSect': round(sum(par_sect.values()) / len(par_sect), 1), 'secteurs': len(par_sect)}
    return {'f': fen, 'couv': couverture_mois(d)}


def couverture_mois(d, plancher=5):
    """28/10/2026 (Will : « avertir à l'écran » quand un mois est incomplet, jamais un texte en dur) — pour chaque mois après le premier :
    secteurs ATTENDUS = ceux qui avaient au moins `plancher` officines le mois d'avant ; secteur COMPLET = attendu qui a ce mois-ci au moins
    la moitié de ses officines du mois d'avant (absent = 0). Rend {'AAAA-MM': [complets, attendus]} pour les seuls mois où complets < attendus.
    Aucun prénom ni nombre d'officines par secteur : ce champ est recopié dans le jeu de chaque commercial."""
    mois = d['mois']
    K = [int(m[:4] + m[5:7]) for m in mois]
    par = {}
    for r in d['r']:
        par.setdefault(r[2], {}).setdefault(r[0], set()).add(r[1])
    out = {}
    for i in range(1, len(K)):
        avant, ici = par.get(K[i - 1], {}), par.get(K[i], {})
        attendus = [c for c, s in avant.items() if len(s) >= plancher]
        complets = [c for c in attendus if 2 * len(ici.get(c, ())) >= len(avant[c])]
        if len(complets) < len(attendus):
            out[mois[i]] = [len(complets), len(attendus)]
    return out


def lire_wml_v2(chemin):
    """SOURCES, RATTACHEES, avoirs et mois lus (couples (année, mois)) DANS generate_wml_v2.py (non importable) : jamais de copie qui diverge."""
    src = open(chemin, encoding='utf-8').read()
    out = {}
    for node in ast.parse(src).body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name) \
                and node.targets[0].id in SOURCES_ENV:
            if node.targets[0].id == 'MOIS_LUS':   # une compréhension de liste : pas un littéral, on l'évalue sans rien d'autre que range
                out['MOIS_LUS'] = eval(compile(ast.Expression(node.value), chemin, 'eval'), {'__builtins__': {'range': range}})
            else:
                out[node.targets[0].id] = ast.literal_eval(node.value)
    manque = [k for k in SOURCES_ENV if k not in out]
    if manque:
        raise ValueError('generate_wml_v2.py : %s introuvable' % ', '.join(manque))
    return out


def lire_js_liste(chemin, variable, cle=None):
    """Liste de CIP13 d'un `window.<variable> = ...;` (liste directe, ou objet {cle:[...]})."""
    t = '\n'.join(l for l in open(chemin, encoding='utf-8').read().split('\n') if not l.lstrip().startswith('//'))   # commentaires de ligne
    t = t[t.index('window.' + variable):]
    t = t[t.index('=') + 1:].strip()
    debut = t.index('{' if cle else '[')
    v = json.JSONDecoder().raw_decode(t[debut:])[0]
    return v[cle] if cle else v


def premier_existant(*chemins):
    for c in chemins:
        if c and os.path.exists(c):
            return c
    return None


def lire_catalogue(stats):
    """Catalogue produit : CIP13 -> (sous-famille, nature, laboratoire, afmcode). Source : STATS/stock POS.xlsx, puis
    le plus ancien tarif en complément. Valeurs #N/A ou vides = None."""
    import openpyxl
    cat = {}
    sources = [os.path.join(stats, 'stock POS.xlsx'), os.path.join(stats, 'stock et prix 22 06 2026.xlsx')]
    utilises = []
    for p in sources:
        if not os.path.exists(p):
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        it = wb.active.iter_rows(values_only=True)
        h = [str(x).strip().lower() if x is not None else '' for x in next(it)]
        need = ('artcodebarre', 'artsousfamille', 'artnature', 'artcollection', 'afmcode')
        if any(n not in h for n in need):
            wb.close()
            continue
        ix = [h.index(n) for n in need]
        n0 = len(cat)
        for r in it:
            c = cip13(r[ix[0]] if ix[0] < len(r) else None)
            if not c or c in cat:
                continue
            vals = []
            for i in ix[1:]:
                v = r[i] if i < len(r) else None
                vals.append(None if v is None or str(v).strip() in ('', '#N/A') else str(v).strip())
            cat[c] = tuple(vals)
        wb.close()
        utilises.append('%s (+%d)' % (os.path.basename(p), len(cat) - n0))
    return cat, utilises


def signature_fichiers(*chemins):
    return '|'.join('%s:%s:%d' % (os.path.basename(c), os.stat(c).st_size, int(os.stat(c).st_mtime)) for c in chemins if c and os.path.exists(c))


# ══════════════════════════════════════════ lecture des tableurs (workers) ═══════════════════════
_G = {}


def _init_worker(chemin_cat, labos, avoirs):
    with open(chemin_cat, encoding='utf-8') as f:
        _G['cat'] = {k: tuple(v) for k, v in json.load(f).items()}
    _G['labos'] = tuple(labos)
    _G['avoirs'] = set(avoirs)


def lire_fichier(tache):
    """Un tableur -> fichier de cache {g: groupes de lignes, j: couples (officine, date) des factures FC}, JSON compressé.
    Rend (chemin du cache, nb de lignes lues)."""
    comm, prefix, chemin, garder, ym, sig = tache
    st = os.stat(chemin)
    nom = '%s_%d_%s_%d_%d_%s.json.gz' % (prefix, ym, hashlib.md5((chemin + '|' + ','.join(garder or [])).encode()).hexdigest()[:8],
                                          st.st_size, int(st.st_mtime), sig)
    fc = os.path.join(CACHE, nom)
    if os.path.exists(fc):
        return fc, -1
    import openpyxl
    from pont_officines import code_officine
    cat, labos, avoirs = _G['cat'], _G['labos'], _G['avoirs']
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    it = wb.active.iter_rows(values_only=True)
    h = [str(x).strip() if x is not None else '' for x in next(it)]
    hi = {k: i for i, k in enumerate(h)}

    def c(r, k):
        i = hi.get(k)
        return r[i] if (i is not None and i < len(r)) else None
    G = {}
    jours = set()
    n = 0
    for r in it:
        tir = c(r, 'TIRCODE')
        if tir is None:
            continue
        try:
            code = str(int(float(tir)))
        except (TypeError, ValueError):
            code = str(tir).strip()
        code, _ancien = code_officine(code)      # comme generate_wml_v2 : une pharmacie = une fiche
        if not code or (garder is not None and code not in garder):
            continue
        q, B, N, net = num(c(r, 'PLVQTE')), num(c(r, 'PLVPUBRUT')), num(c(r, 'PLVPUNET')), num(c(r, 'PLVMNTNETHT'))
        pcv = str(c(r, 'PCVNUM') or '')
        if prefix in avoirs and pcv.startswith('AC'):
            q, net = -q, -net                     # avoirs Escale, comme generate_wml_v2
        if pcv.startswith('FC') and code != 'INTER':   # jour de commande = une date avec au moins une facture (page de réunion du 28/10/2026)
            dt = c(r, 'PCVDATEEFFET')
            jours.add((code, dt.date().isoformat() if hasattr(dt, 'date') else str(dt)[:10]))
        cip = cip13(c(r, 'ARTCODEBARRE'))
        k = cat.get(cip) or (None, None, None, None)
        afm = c(r, 'AFMCODE') if 'AFMCODE' in hi else k[3]
        t = c(r, 'TVATAUX')
        t = round(t, 4) if isinstance(t, (int, float)) else None
        remb = str(afm or '').strip() == 'REMBSS'
        rb = remb and (t is None or abs(t - 0.021) < 1e-6)
        froid = sa(c(r, 'ARTSOUSFAMILLE') or k[0]) == 'froid'
        gen = 'generique' in sa(c(r, 'ARTNATURE') or k[1])
        lab = str(c(r, 'ARTCOLLECTION') or k[2] or '').upper()
        su = any(x in lab for x in labos)
        mt = c(r, 'MARGETHEORIQUE')
        a_mt = isinstance(mt, (int, float))
        reg = regime_ligne(c(r, 'PCVDATEEFFET'), ym)
        d = str(c(r, 'PLVDESIGNATION') or '').strip()
        cle = cip or '#' + d[:40]
        key = (code, reg, cle, round(B, 4), round(N, 4), int(rb), int(froid), int(gen), int(su), int(q >= 0), int(a_mt))
        x = G.get(key)
        if x is None:
            x = G[key] = [0.0, 0.0, 0.0, 0.0]
        x[0] += q
        x[1] += net
        x[2] += q * N
        if a_mt:
            x[3] += mt
        n += 1
    wb.close()
    out = [list(k) + [round(v, 6) for v in x] for k, x in G.items()]
    tmp = fc + '.tmp'
    with gzip.open(tmp, 'wt', encoding='utf-8', compresslevel=1) as f:
        json.dump({'g': out, 'j': sorted(jours)}, f, separators=(',', ':'))
    os.replace(tmp, fc)
    return fc, n


def charger_jours(chemin):
    """Couples [officine, 'AAAA-MM-JJ'] des factures FC d'un fichier de cache."""
    with gzip.open(chemin, 'rt', encoding='utf-8') as f:
        return json.load(f)['j']


def charger_groupes(chemin):
    """Groupes d'un fichier de cache -> tuples (off, reg, cle, B, N, rb, froid, gen, su, pos, a_mt, q, net, qN, mt)."""
    with gzip.open(chemin, 'rt', encoding='utf-8') as f:
        rows = json.load(f)['g']
    return [(r[0], r[1], r[2], r[3], r[4], bool(r[5]), bool(r[6]), bool(r[7]), bool(r[8]), bool(r[9]), bool(r[10]),
             r[11], r[12], r[13], r[14]) for r in rows]


# ═══════════════════════════════════════════════ principal ═══════════════════════════════════════
def main():
    t0 = time.time()
    prive = lire_prive()
    if prive is None:
        print('ARRÊT : réglage privé illisible (%s) — rien n\'est écrit.' % PRIVE)
        return 3
    try:
        env = lire_wml_v2(os.path.join(BASE, 'generate_wml_v2.py'))
    except (OSError, ValueError, SyntaxError) as e:
        print('ARRÊT :', e)
        return 2
    stats = premier_existant(os.path.join(BASE, 'STATS'), os.path.expanduser('~/JARVIS/APP-atelier/STATS'))
    crm = os.path.join(BASE, 'crm')
    p_gen = premier_existant(os.path.join(crm, 'v2', 'marge-regles-data.js'))
    p_fr = premier_existant(os.path.join(crm, 'v2', 'froid-data.js'))
    p_pr = premier_existant(os.path.join(crm, 'v2', 'marge-regles-prive.js'), os.path.expanduser('~/JARVIS/APP-atelier/crm/v2/marge-regles-prive.js'))
    manque = [n for n, p in (('STATS', stats), ('marge-regles-data.js', p_gen), ('froid-data.js', p_fr)) if not p]
    if manque:
        print('ARRÊT : introuvable :', ', '.join(manque))
        return 2
    gen = lire_js_liste(p_gen, 'MARGE_REGLES', 'generiques')
    fr = lire_js_liste(p_fr, 'FROID_CIPS')
    nr = set()   # non remboursé = critère de la ligne (rb), voir famille()
    privee = []
    if p_pr:
        try:
            privee = (lire_js_liste(p_pr, 'MARGE_REGLES_PRIVE', 'petitReduit') or {}).get('cips') or []
        except (ValueError, KeyError, AttributeError):
            privee = []
    if not (gen and fr):
        print('ARRÊT : une liste de produits est vide (génériques %d · froid %d)' % (len(gen), len(fr)))
        return 2
    L = Listes(gen, nr, fr, privee, prive['petit'], prive['petit_2027'], prive['labos'])
    print('Listes : %d génériques · %d froid · liste privée %s' % (len(L.gen), len(L.fr), ('%d' % len(L.privee)) if privee else 'absente'))

    # ── quels tableurs : exactement ceux de generate_wml_v2 ──
    mois_lus = env['MOIS_LUS']
    rattachees = {k: tuple(v) for k, v in env['RATTACHEES'].items()}
    lectures = [(comm, prefix, stats, None) for comm, prefix in env['SOURCES']]
    for comm, prefix in sorted(set(rattachees.values())):
        lectures.append((comm, prefix, os.path.join(stats, 'hors-perimetre'), sorted(k for k, v in rattachees.items() if v == (comm, prefix))))
    cat, utilises = lire_catalogue(stats)
    if not cat:
        print('ARRÊT : catalogue produit illisible dans', stats)
        return 2
    print('Catalogue : %d produits (%s)' % (len(cat), ' ; '.join(utilises)))
    os.makedirs(CACHE, exist_ok=True)
    chemin_cat = os.path.join(CACHE, 'catalogue.json')
    sig_cat = signature_fichiers(os.path.join(stats, 'stock POS.xlsx'), os.path.join(stats, 'stock et prix 22 06 2026.xlsx'))
    sig = hashlib.md5((VERSION_CACHE + sig_cat + ' '.join(utilises) + '|'.join(L.labos) + '|' + str(DEBUT_2027) + str(MOIS_2027)).encode()).hexdigest()[:8]
    with open(chemin_cat, 'w', encoding='utf-8') as f:
        json.dump(cat, f, ensure_ascii=False)
    taches = []
    for comm, prefix, dossier, garder in lectures:
        for annee, mois in mois_lus:
            p = os.path.join(dossier, '%s_%02d_%d.xlsx' % (prefix, mois, annee))
            if garder is not None and not os.path.exists(p):
                p = os.path.join(stats, os.path.basename(p))
            if os.path.exists(p):
                taches.append((comm, prefix, p, garder, annee * 100 + mois, sig))
    if not taches:
        print('ARRÊT : aucun tableur de ventes lu dans', stats)
        return 2
    print('%d tableurs · cache %s' % (len(taches), CACHE), flush=True)
    chemins = []
    nouveaux = 0
    with Pool(4, initializer=_init_worker, initargs=(chemin_cat, L.labos, env['AVOIRS_EN_POSITIF'])) as pool:
        for i, (fc, n) in enumerate(pool.imap(lire_fichier, taches), 1):
            chemins.append(fc)
            nouveaux += n >= 0
            if i % 25 == 0:
                print('  %d/%d (%.0f s)' % (i, len(taches), time.time() - t0), flush=True)
    meta = [(t[0], t[4]) for t in taches]       # (prénom, ym) de chaque fichier, dans l'ordre
    print('Lecture : %d fichiers lus pour de bon, %d repris du cache (%.0f s)' % (nouveaux, len(taches) - nouveaux, time.time() - t0), flush=True)

    # ── passe 1 : apprendre les prix d'achat ; passe 2 : estimer ──
    cascade = Cascade(L)
    nb_lignes = 0
    ventes = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))     # (régime, code) -> ym -> brut -> boîtes
    for fc, (comm, ym) in zip(chemins, meta):
        for g in charger_groupes(fc):
            cascade.apprendre(g, ym)
            if g[9] and g[11] > 0 and g[2] not in L.gen:
                ventes[(g[1], g[2])][ym][g[3]] += g[11]
            nb_lignes += 1
    if not cascade.figer():
        print('ARRÊT : aucune ligne avec la marge de la direction (MARGETHEORIQUE) — rien à apprendre.')
        return 4
    print('Apprentissage : %d groupes · %d portent la marge de la direction · %d prix produit-mois-brut · %d produit-brut · %d produit (%.0f s)'
          % (nb_lignes, cascade.lignes_connues, len(cascade.OM), len(cascade.OB), len(cascade.OC), time.time() - t0), flush=True)
    lignes = {}
    prov = defaultdict(float)
    rempl = prive['remplacements']
    baisses = detecter_baisses(cascade.OM, ventes, {ym: set(x['cips']) for ym, x in rempl.items()})
    suivis = defaultdict(dict)                      # mois de remplacement ou de référence -> {code: produit}
    for x in rempl.values():
        suivis[x['par']].update(x['cips'])
    for d_b in baisses.values():
        for cle_b, (n_b, _t, _bp) in d_b.items():
            if n_b is not None:
                suivis[n_b].setdefault(cle_b, cle_b)    # baisse : le « produit » est le code lui-même
    tables = {ym: ({}, {}) for ym in suivis}        # mois de remplacement -> marge et boîtes par code / produit, par brut
    differes, differes_b = [], []
    for fc, (comm, ym) in zip(chemins, meta):
        r = rempl.get(ym)
        bx = baisses.get(ym)
        for g in charger_groupes(fc):
            if r and g[2] in r['cips']:
                differes.append((comm, ym, g))      # estimé après la lecture du mois de remplacement
                continue
            if bx and g[2] in bx and g[3] < bx[g[2]][2] * SEUIL_BAISSE:
                differes_b.append((comm, ym, g))    # estimé après la lecture du mois de référence
                continue
            src, m = ajouter(lignes, (comm, g[0], ym), g, ym, cascade, L)
            if src != 'gen':
                prov[src] += m
                nom = suivis[ym].get(g[2]) if ym in suivis else None
                if nom and g[11] > 0:
                    for t, k in ((tables[ym][0], g[2]), (tables[ym][1], nom)):
                        x = t.setdefault(k, {}).setdefault(g[3], [0.0, 0.0])
                        x[0] += m
                        x[1] += g[11]
    bilan = defaultdict(lambda: [0.0, 0.0, 0.0, defaultdict(float)])   # (mois, produit) -> boîtes, marge sans, marge avec, boîtes par niveau
    for comm, ym, g in differes:
        r = rempl[ym]
        nom = r['cips'][g[2]]
        pb, niveau = par_boite(tables[r['par']], g[2], nom, g[3])
        sans = cascade.marge(g, ym)[0]
        src, m = ajouter(lignes, (comm, g[0], ym), g, ym, cascade, L, None if pb is None else g[11] * pb)
        if src != 'gen':
            prov[src] += m
        b = bilan[(ym, nom)]
        b[0] += g[11]
        b[1] += sans
        b[2] += m
        b[3][niveau or 'non vendu : cascade ordinaire'] += g[11]
    for (ym, nom), b in sorted(bilan.items()):
        print('Remplacement %d-%02d par %d-%02d · %s : %.0f boîtes · marge %.0f € -> %.0f € (%s)'
              % (ym // 100, ym % 100, rempl[ym]['par'] // 100, rempl[ym]['par'] % 100, nom, b[0], b[1], b[2],
                 ' · '.join('%s %.0f' % kv for kv in sorted(b[3].items()))))
    bilan_b = defaultdict(lambda: [set(), 0.0, 0.0, 0.0, defaultdict(float)])   # mois -> codes, boîtes, marge sans, avec, boîtes par niveau
    for comm, ym, g in differes_b:
        sans = cascade.marge(g, ym)[0]
        imp, niveau = marge_baisse(g, baisses[ym][g[2]], tables, sans)
        src, m = ajouter(lignes, (comm, g[0], ym), g, ym, cascade, L, imp)
        if src != 'gen':
            prov['baisse de prix' if src == 'mois de remplacement' else src] += m
        b = bilan_b[ym]
        b[0].add(g[2])
        b[1] += g[11]
        b[2] += sans
        b[3] += m
        b[4][niveau or 'non vendu : cascade ordinaire'] += g[11]
    for ym, b in sorted(bilan_b.items()):
        print('Baisses de prix %d-%02d : %d codes · %.0f boîtes · marge %.0f € -> %.0f € (%s)'
              % (ym // 100, ym % 100, len(b[0]), b[1], b[2], b[3], ' · '.join('%s %.0f' % kv for kv in sorted(b[4].items()))))
    reserves = dict(prive['reserves'])
    for ym, d_b in sorted(baisses.items()):
        k = '%d-%02d' % (ym // 100, ym % 100)
        t_b = texte_baisses(d_b, ym, autres=ym in rempl)
        reserves[k] = reserves[k] + ' ; ' + t_b if ym in rempl else t_b
    prenoms = sorted({k[0] for k in lignes})
    officines = sorted({k[1] for k in lignes})
    c_idx = {p: i for i, p in enumerate(prenoms)}
    o_idx = {o: i for i, o in enumerate(officines)}
    rows = arrondir(lignes, c_idx, o_idx)
    if not rows:
        print('ARRÊT : 0 ligne — rien n\'est écrit.')
        return 5
    # ── jours de commande : dates distinctes par officine et par mois, tous commerciaux et tous tableurs confondus ──
    J = jours_par_mois(x for fc in chemins for x in charger_jours(fc))
    for r in rows:
        r.append(J.get((officines[r[1]], r[2]), 0))
    vus = {(officines[r[1]], r[2]): r[13] for r in rows}      # ce que la fiche voit : les jours des officine-mois qui ont une ligne
    print('Jours de commande : %d jours · %d officines · dont %d jours sur des officine-mois présents dans le fichier'
          % (sum(J.values()), len({k[0] for k in J}), sum(vus.values())))
    mois_txt = ['%d-%02d' % (a, m) for a, m in mois_lus]
    ym_max = max(m[1] for m in meta)
    d = {'v': 1, 'fait': datetime.date.today().isoformat(), 'jusqua': '%d-%02d' % (ym_max // 100, ym_max % 100), 'mois': mois_txt,
         'reserves': reserves, 'fam': FAM, 'c': prenoms, 'o': officines, 'r': rows}
    cout = prive['cout_livraison']
    moy = moyenne_livraison(d, vus, cout) if cout else None
    if moy:
        d['livraison'] = {'cout': cout, 'moy': moy}
        print('Livraison : coût %.2f € · officine moyenne de %d clients : %.2f jours de commande par mois · %.2f € par jour · %.2f € par mois livraisons déduites'
              % (cout, moy['n'], moy['jm'], moy['parJour'], moy['renta']))
        # Moyenne du Pilotage par fenêtre : la fenêtre entière DOIT retomber sur la moyenne de la fiche, sinon rien n'est écrit.
        pm = moyennes_pilotage(d, cout)
        tout = pm['f'].get(mois_txt[0] + '|' + mois_txt[-1])
        ecarts = [q for q in ('n', 'jm', 'parJour', 'renta') if tout is None or abs(tout[q] - moy[q]) > 0.001]
        if ecarts:
            print('ARRÊT : la fenêtre entière du Pilotage ne retombe pas sur livraison.moy (%s) — rien n\'est écrit.' % ', '.join(ecarts))
            return 5
        d['pilotageMoy'] = pm
        print('Pilotage : %d fenêtres de mois, la fenêtre entière retombe sur la moyenne de la fiche.' % len(pm['f']))
    else:
        print('Livraison : coût absent du réglage privé (cout_livraison) — pas de bloc « Rentabilité » sur la fiche.')
    for r in rows:
        assert abs(sum(r[5:12]) - r[4]) < 0.005, 'tot != somme des familles'

    # ── écriture atomique + relecture + node --check ──
    texte = formater(d)
    tmp = os.path.join(BASE, 'marge-clients.nouveau.js')
    try:
        with open(tmp, 'w', encoding='utf-8') as f:
            f.write(texte)
        if subprocess.run(['node', '--check', tmp], capture_output=True).returncode != 0:
            print('ARRÊT : node --check refuse le fichier — rien n\'est remplacé.')
            return 5
        rel = analyser(open(tmp, encoding='utf-8').read())
        if rel != d:
            print('ARRÊT : la relecture ne redonne pas ce qui a été écrit — rien n\'est remplacé.')
            return 5
        os.replace(tmp, SORTIE)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    # cache : on garde seulement ce que cette passe a utilisé
    gardes = set(chemins) | {chemin_cat}
    for f in glob.glob(os.path.join(CACHE, '*.json.gz')) + glob.glob(os.path.join(CACHE, '*.tmp')):
        if f not in gardes:
            os.remove(f)
    tp = sum(prov.values())
    print('OK %s · %d lignes · %d clients (prénoms) · %d officines · %.1f Mo · %.0f s'
          % (SORTIE, len(rows), len(prenoms), len(officines), os.path.getsize(SORTIE) / 1e6, time.time() - t0))
    print('Provenance de la marge : ' + ' · '.join('%s %.3f %%' % (k, 100 * v / tp) for k, v in sorted(prov.items(), key=lambda kv: -kv[1])) if tp else 'Provenance : rien')
    return 0


if __name__ == '__main__':
    sys.exit(main())
