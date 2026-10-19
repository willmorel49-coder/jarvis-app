# -*- coding: utf-8 -*-
"""CATALOGUE DU GROUPE (7 établissements) -> crm/v2/catalogue-groupe-data.js

Alimente la création de fiche Marketing : UN catalogue, remboursables et non remboursables
mêlés, SANS Offilog. Rejouable : /usr/bin/python3 generate_catalogue_groupe.py (Python 3.9,
openpyxl + xlrd). Lit des fichiers locaux HORS dépôt (exports des établissements, pont de codes).

SOURCES (reconnues à leurs COLONNES ; seul le site se lit dans le nom du fichier) :
  · ~/JARVIS/PRIX ET STOCKS ETABLISSEMENTS/ : export complet POS (avec PPHT), export complet SEP
    (stock seul), extraits « non remboursables » des 7 sites, fichier STOCK .xls (5 sites)
  · ~/JARVIS/APP/STATS/ : ancien catalogue de juin (molécules, tarif de repli), pont de codes
  · dépôt : prod-stats-data.js (pharmacies acheteuses, famille réseau), generiqueurs-data.js,
    mkt-rayons-data.js + saison-cip.json + marketing-mix-data.js (rayons)

RÈGLES ARRÊTÉES LE 19/10/2026 :
  · Tarif = PPHT, LE MÊME pour les 7 établissements : celui de l'export POS. Replis datés
    (autre export complet, catalogue de juin, extraits NR) SEULEMENT quand POS n'a pas de tarif.
  · La colonne `atfprix` n'est JAMAIS lue comme tarif. Aucune colonne de prix d'achat ou de revient
    (ARTLASTPA, artprmp, atfprix) n'est lue, donc aucune ne peut sortir.
  · Aucun produit ni aucune donnée Offilog.
  · `fam` et `net` : MÊMES règles que generate_catalogue_complet.py, dont les fonctions sont
    importées (famille, abandon_ip, porte_abandon, labos_generiqueurs, r2), pas recopiées.
  · La MARQUE n'est jamais un fournisseur / dépositaire : voir resoudre_marque().
"""
import contextlib
import csv
import glob
import io
import json
import os
import re
import sys
import time
import unicodedata
from collections import Counter, defaultdict

import openpyxl
import xlrd

ROOT = os.path.dirname(os.path.abspath(__file__))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)
V2 = os.path.join(ROOT, 'crm', 'v2')
HOME = os.path.expanduser('~')
SRC = os.path.join(HOME, 'JARVIS', 'PRIX ET STOCKS ETABLISSEMENTS')
STATS = os.path.join(HOME, 'JARVIS', 'APP', 'STATS')
OUT = os.path.join(V2, 'catalogue-groupe-data.js')
SITES = ['OPS', 'HP', 'CPR', 'POS', 'SEP', 'MSP', 'SOP']
SITE_BASE_XLS = 'SOP'   # le fichier STOCK .xls liste les produits en stock chez SOP (stocklivrablesop > 0 partout)
XLS_COLS = {'stocklivrablesop': 'SOP', 'stockmsp': 'MSP', 'stockhp': 'HP', 'stockcpr': 'CPR', 'stockops': 'OPS'}
SITE_PRIO = ['POS', 'OPS', 'CPR', 'MSP', 'SOP', 'HP', 'SEP']   # départage des tarifs d'extraits
FORBIDDEN = ('artlastpa', 'artprmp', 'atfprix', 'prix_achat', 'revient')


def log(*a):
    print(*a, flush=True)


def importer_regles():
    """Importe generate_catalogue_complet (fam, net) sans que son lecteur de tarif (qui exige un
    dossier STATS/ du dépôt) ne s'exécute : seul le choix de fichier est neutralisé, aucune règle."""
    import source_prix
    source_prix.source_prix = lambda *a, **k: ''
    with contextlib.redirect_stdout(io.StringIO()):
        import generate_catalogue_complet as gcc
    return gcc


# ───────────────────────── utilitaires ─────────────────────────
def code_barres(v):
    s = re.sub(r'\.0+$', '', str(v if v is not None else '').strip())
    return s if re.fullmatch(r'\d{12,14}', s) else ''


def cip13_of(c7):
    b = '34009' + c7
    s = sum(int(d) * (1 if i % 2 == 0 else 3) for i, d in enumerate(b))
    return b + str((10 - s % 10) % 10)


def nombre(v):
    try:
        return float(str(v).replace(',', '.'))
    except (TypeError, ValueError):
        return 0.0


def txt(v):
    if v is None:
        return ''
    s = str(v).strip()
    return '' if s in ('#N/A', 'None', 'nan') else s


def norm(s):
    """MAJUSCULES sans accent ni ponctuation, mots séparés par une espace."""
    s = unicodedata.normalize('NFD', (s or '').upper())
    s = ''.join(ch for ch in s if unicodedata.category(ch) != 'Mn')
    return re.sub(r'[^A-Z0-9]+', ' ', s).strip()


def date_nom(nom, mtime):
    m = re.search(r'(\d{2}) (\d{2}) (\d{4})', nom)
    if m:
        return '%s-%s-%s' % (m.group(3), m.group(2), m.group(1))
    m = re.search(r'(?<!\d)(\d{2})(\d{2})(20\d{2})(?!\d)', nom)
    if m:
        return '%s-%s-%s' % (m.group(3), m.group(2), m.group(1))
    m = re.search(r'(?<!\d)(\d{2})(\d{2})(\d{2})(?!\d)', nom)
    if m and 1 <= int(m.group(2)) <= 12:
        return '20%s-%s-%s' % (m.group(3), m.group(2), m.group(1))
    return time.strftime('%Y-%m-%d', time.localtime(mtime))


def lire_xlsx(chemin):
    ws = openpyxl.load_workbook(chemin, read_only=True, data_only=True).active
    it = ws.iter_rows(values_only=True)
    h = [str(x or '').strip().lower() for x in next(it)]
    return h, it


def entete_xlsx(chemin):
    try:
        ws = openpyxl.load_workbook(chemin, read_only=True, data_only=True).active
        return [str(c.value or '').strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    except Exception:
        return []


def js_json(chemin, ouvrant='{', fermant='}'):
    t = open(chemin, encoding='utf-8').read()
    return json.loads(t[t.index(ouvrant):t.rindex(fermant) + 1])


# ───────────────────────── découverte des sources par COLONNES ─────────────────────────
def decouvrir():
    out = {'complet_ppht': {}, 'complet_sans_ppht': {}, 'extraits': {}, 'xls': None, 'juin': None}
    vus = set()
    for f in sorted(glob.glob(os.path.join(SRC, '*.xlsx'))) + sorted(glob.glob(os.path.join(STATS, 'stock*.xlsx'))):
        b = os.path.basename(f)
        if b.startswith('~$'):
            continue
        h = set(entete_xlsx(f))
        mt = os.path.getmtime(f)
        m = re.match(r'\s*(?:stock\s+)?([A-Za-z]{2,4})\b', b)
        site = m.group(1).upper() if m else None
        if {'artlastpa', 'ppht', 'stockdispo', 'artcodebarre'} <= h and 'artcode' not in h:
            out['extraits'][site] = {'f': f, 'date': time.strftime('%Y-%m-%d', time.localtime(mt))}
        elif {'artcodebarre', 'ppht', 'stockdispo', 'afmcode', 'artcode'} <= h and 'artmol' not in h:
            cle = (os.path.getsize(f), site)
            if cle in vus:
                continue
            vus.add(cle)
            d = date_nom(b, mt)
            cur = out['complet_ppht'].get(site)
            if not cur or d > cur['date']:
                out['complet_ppht'][site] = {'f': f, 'date': d}
        elif {'artcodebarre', 'stockdispo', 'afmcode', 'artcode'} <= h and 'ppht' not in h:
            d = date_nom(b, mt)
            cur = out['complet_sans_ppht'].get(site)
            if not cur or d > cur['date']:
                out['complet_sans_ppht'][site] = {'f': f, 'date': d}
        elif {'artcodebarre', 'ppht', 'artmol', 'afmcode'} <= h:
            d = date_nom(b, mt)
            if not out['juin'] or d > out['juin']['date']:
                out['juin'] = {'f': f, 'date': d}
    for f in glob.glob(os.path.join(SRC, '*.xls')):
        try:
            s = xlrd.open_workbook(f).sheet_by_index(0)
            h = [str(c).strip().lower() for c in s.row_values(0)]
        except Exception:
            continue
        if 'artcode' in h and any(c in h for c in XLS_COLS):
            out['xls'] = {'f': f, 'date': date_nom(os.path.basename(f), os.path.getmtime(f))}
    return out


# ───────────────────────── pont de codes ─────────────────────────
ALIAS = {}
LIBELLES = {}
RP_RE = re.compile(r'\s*-?\s*\bRP\s*\d{7}\b')


def charger_pont():
    p = os.path.join(STATS, 'pont-codes.json')
    try:
        brut = json.load(open(p, encoding='utf-8'))['alias']
        # un alias qui ne rend pas un vrai code-barres (ex. 20 chiffres collés : TUSSONYL) est ignoré
        ALIAS.update({k: v for k, v in brut.items() if code_barres(v)})
        log('pont de codes : %d alias écartés (valeur qui n\'est pas un code-barres)' % (len(brut) - len(ALIAS)))
        log('pont de codes : %d alias (%s)' % (len(ALIAS), time.strftime('%Y-%m-%d', time.localtime(os.path.getmtime(p)))))
    except (OSError, ValueError, KeyError):
        log('ATTENTION pont-codes.json illisible : aucun alias appliqué')
    try:
        for r in csv.DictReader(open(os.path.join(ROOT, 'fusions-codes.csv'), encoding='utf-8'), delimiter=';'):
            g = r['garde_code_barres'] or r['garde_cip13']
            if r['libelle']:
                LIBELLES[ALIAS.get(g, g)] = r['libelle']
    except OSError:
        pass


def canon(c):
    return ALIAS.get(c, c)


# ───────────────────────── fournisseurs : jamais une marque ─────────────────────────
FOURNISSEUR_RX = re.compile(r'^(INTER TRADE|ITP$|ITP )')
FOURNISSEURS_VUS = set()    # toutes les graphies lues dans les fichiers (pour meta.fournisseurs)


def est_fournisseur(labo):
    return bool(labo) and bool(FOURNISSEUR_RX.match(norm(labo)))


def voir_labo(v):
    """Rend le texte du labo ; mémorise la graphie si c'est un fournisseur."""
    t = txt(v)
    if t and est_fournisseur(t):
        FOURNISSEURS_VUS.add(t)
    return t


# ───────────────────────── lecture des sources ─────────────────────────
ANOM = Counter()


def lire_complet(f, avec_ppht, nom):
    """Export complet d'un site. Ne garde QUE : code, désignation, labo(s), afm, nature, ppht, stock.
    `artmarque` (SEP) est gardé à part : c'est une AUTRE colonne qui porte parfois la vraie marque."""
    h, it = lire_xlsx(f)
    ix = {n: i for i, n in enumerate(h)}
    ci, di, ai = ix['artcodebarre'], ix['artdesignation'], ix['afmcode']
    li = ix.get('artcollection')
    mqi = ix.get('artmarque')
    ni = ix.get('artnature')
    pi = ix['ppht'] if avec_ppht else None
    si, ki = ix['stockdispo'], ix['artcode']
    out, cip7 = {}, {}
    for r in it:
        brut = code_barres(r[ci])
        lib = txt(r[di])
        a7_ = str(r[ki] or '').strip()
        if not brut and re.fullmatch(r'\d{7}', a7_):
            brut = cip13_of(a7_)
            ANOM[nom + ' : CIP13 recalculé depuis artcode'] += 1
        if not brut:
            ANOM[nom + ' : ligne sans code-barres ni CIP7 (écartée)'] += 1
            continue
        c = canon(brut)
        try:
            stock = max(0, int(float(r[si] or 0)))
        except (TypeError, ValueError):
            stock = 0
        p = round(nombre(r[pi]) + 1e-9, 2) if pi is not None else 0.0
        a7 = str(r[ki] or '').strip()
        if len(a7) == 7 and a7.isdigit():
            cip7.setdefault(a7, c)
        labo = voir_labo(r[li]) if li is not None else ''
        marque = voir_labo(r[mqi]) if mqi is not None else ''
        o = out.get(c)
        if o is None:
            out[c] = {'d': lib, 'labo': labo, 'marque': marque, 'afm': txt(r[ai]),
                      'nat': txt(r[ni]) if ni is not None else '', 'ppht': p, 's': stock}
        else:
            ANOM[nom + ' : code-barres en double (stocks additionnés)'] += 1
            o['s'] += stock
            if o['ppht'] <= 0 and p > 0:
                o['ppht'] = p
            if not o['d']:
                o['d'] = lib
            o['labo'] = o['labo'] or labo
            o['marque'] = o['marque'] or marque
    return out, cip7


def lire_extrait(f, site):
    h, it = lire_xlsx(f)
    ix = {n: i for i, n in enumerate(h)}
    ci, di, li, ai, pi, si = (ix['artcodebarre'], ix['artdesignation'], ix['artcollection'],
                              ix['afmcode'], ix['ppht'], ix['stockdispo'])
    out = {}
    for r in it:
        brut = code_barres(r[ci])
        if not brut:
            ANOM['extrait %s : ligne sans code-barres (écartée)' % site] += 1
            continue
        c = canon(brut)
        p = round(nombre(r[pi]) + 1e-9, 2)
        try:
            s = max(0, int(float(r[si] or 0)))
        except (TypeError, ValueError):
            s = 0
        o = out.setdefault(c, {'d': txt(r[di]), 'labo': voir_labo(r[li]), 'afm': txt(r[ai]), 'ppht': 0.0, 's': 0})
        o['s'] += s
        if o['ppht'] <= 0 and p > 0:
            o['ppht'] = p
    return out


def lire_xls(f, cip7_map):
    """Fichier STOCK : on ne lit QUE le code, la désignation et les colonnes de stock (atfprix ignoré)."""
    s = xlrd.open_workbook(f).sheet_by_index(0)
    h = [str(c).strip().lower() for c in s.row_values(0)]
    ia = h.index('artcode')
    idg = h.index('artdesignation') if 'artdesignation' in h else None
    cols = {XLS_COLS[c]: h.index(c) for c in XLS_COLS if c in h}
    out = {}
    for k in range(1, s.nrows):
        r = s.row_values(k)
        a = str(r[ia]).strip()
        if re.fullmatch(r'\d{7}', a):
            c = cip7_map[a] if a in cip7_map else canon(cip13_of(a))
        elif code_barres(a):
            c = canon(code_barres(a))
        else:
            ANOM['xls : artcode inexploitable'] += 1
            continue
        e = out.setdefault(c, {'d': txt(r[idg]) if idg is not None else '', 's': {}})
        for site, i in cols.items():
            try:
                v = max(0, int(float(r[i] or 0)))
            except (TypeError, ValueError):
                v = 0
            e['s'][site] = e['s'].get(site, 0) + v
    return out


def lire_juin(f):
    h, it = lire_xlsx(f)
    ix = {n: i for i, n in enumerate(h)}
    out = {}
    for r in it:
        c = code_barres(r[ix['artcodebarre']])
        if not c:
            continue
        c = canon(c)
        lib = txt(r[ix['artdesignation']])
        if not lib:
            continue
        p = round(nombre(r[ix['ppht']]) + 1e-9, 2)
        o = out.get(c)
        if o is None:
            out[c] = {'d': lib, 'labo': voir_labo(r[ix['artcollection']]), 'mol': txt(r[ix['artmol']]),
                      'nat': txt(r[ix['artnature']]), 'afm': txt(r[ix['afmcode']]), 'ppht': p}
        elif o['ppht'] <= 0 and p > 0:
            o['ppht'] = p
    return out


# ───────────────────────── rayons (miroir de generate_mkt_rayons.mjs) ─────────────────────────
def charger_rayons():
    mjs = open(os.path.join(ROOT, 'generate_mkt_rayons.mjs'), encoding='utf-8').read()
    bloc_atc = mjs[mjs.index('const ATC = {'):mjs.index('// Catégories NR')]
    atc = dict(re.findall(r"([A-Z]\d\d):\s*'([^']+)'", bloc_atc))
    bloc_nr = mjs[mjs.index('const NR = {'):mjs.index('const rayons = []')]
    nrmap = dict(re.findall(r"'([^']+)':\s*'([^']+)'", bloc_nr))
    R = js_json(os.path.join(V2, 'mkt-rayons-data.js'))
    connu = {c: R['rayons'][i] for c, i in R['cip'].items()}
    saison = json.load(open(os.path.join(V2, 'saison-cip.json'), encoding='utf-8'))['data']
    mix = js_json(os.path.join(V2, 'marketing-mix-data.js'))
    nr_cip = {}
    for cat in mix['nr']:
        for it in cat['rows']:
            nr_cip.setdefault(it['cip'], nrmap.get(cat['cat'], cat['cat']))
    return connu, atc, saison, nr_cip, list(R['rayons'])


# ───────────────────────── la VRAIE marque ─────────────────────────
# Table de marques CONNUES, pour les produits que les fichiers attribuent à un fournisseur et
# qu'aucun dictionnaire de code-barres ne retrouve. Clé = début de la désignation (norm()).
MARQUES_CONNUES = {
    'AVENE': 'AVENE', 'XERACALM': 'AVENE', 'CICALFATE': 'AVENE',
    'ANTHELIOS': 'LA ROCHE-POSAY', 'LIPIKAR': 'LA ROCHE-POSAY', 'CICAPLAST': 'LA ROCHE-POSAY',
    'POSTHELIOS': 'LA ROCHE-POSAY', 'EFFACLAR': 'LA ROCHE-POSAY',
    'MUSTELA': 'MUSTELA', 'FILORGA': 'FILORGA', 'URIAGE': 'URIAGE', 'SVR': 'SVR', 'BIOLANE': 'BIOLANE',
    'URGO': 'URGO', 'CINQ SUR CINQ': 'CINQ SUR CINQ', 'CINQ CINQ': 'CINQ SUR CINQ',
    'PEDIAKID': 'PEDIAKID', 'GRANIONS': 'GRANIONS', 'GRANIONSKID': 'GRANIONS',
    'DERMOPHIL INDIEN': 'DERMOPHIL INDIEN', 'BAUME DU TIGRE': 'BAUME DU TIGRE', 'DAYANG': 'DAYANG',
    'ACCU CHEK': 'ACCU-CHEK', 'CONTOUR': 'CONTOUR', 'ONETOUCH': 'ONETOUCH', 'FREESTYLE': 'FREESTYLE',
    'MEPILEX': 'MOLNLYCKE', 'ALLEVYN': 'SMITH AND NEPHEW', 'JELONET': 'SMITH AND NEPHEW',
    'AQUACEL': 'CONVATEC', 'CONVAFOAM': 'CONVATEC', 'MICRO FINE': 'EMBECTA', 'BD': 'EMBECTA',
    'SERUM PHYSIO GILBERT': 'GILBERT', 'SERUM PHYSIO GIFRER': 'GIFRER', 'CLEARBLUE': 'CLEARBLUE',
    'AUDISPRAY': 'COOPER', 'CREALINE': 'BIODERMA', 'SEBIUM': 'BIODERMA',
    'COTOPADS': 'COTOPADS', 'MIRADENT': 'MIRADENT', 'TENSOPLAST': 'TENSOPLAST', 'ZETUVIT': 'HARTMANN',
}


class Marques:
    """Dictionnaires pour retrouver la vraie marque d'un produit dont le « labo » est un fournisseur."""

    def __init__(self, prods_reels):
        self.stats = Counter()
        self.canon = {}                 # norm(marque) -> graphie du catalogue la plus fréquente
        c = Counter(l for l in prods_reels if l)
        for l, _ in c.most_common():
            self.canon.setdefault(norm(l), l)
        # code-barres -> marque : dictionnaires de catalogues présents dans le dépôt (le NOM de la
        # marque seulement : aucun prix, aucune autre donnée, rien d'Offilog)
        self.par_code = {}
        for nom, fn in (('labo-cip.json (ventes réseau)', self._labo_cip),
                        ('pharmazon-data.js (marque par code-barres)', self._pharmazon),
                        ('cap3000-data.js (marque par code-barres)', self._cap3000)):
            n = 0
            for ean, m in fn().items():
                if m and not est_fournisseur(m):
                    self.par_code.setdefault(ean, []).append((nom, m))
                    n += 1
            log('  dictionnaire %s : %d codes' % (nom, n))
        self.table = {tuple(norm(k).split()): v for k, v in MARQUES_CONNUES.items()}
        self.longueurs = sorted({len(k) for k in self.table}, reverse=True)

    @staticmethod
    def _labo_cip():
        try:
            return json.load(open(os.path.join(V2, 'labo-cip.json'), encoding='utf-8'))['data']
        except (OSError, ValueError, KeyError):
            return {}

    @staticmethod
    def _pharmazon():
        try:
            t = open(os.path.join(V2, 'pharmazon-data.js'), encoding='utf-8').read()
            return {x['ean']: x.get('labo') for x in json.loads(t[t.index('['):t.rindex(']') + 1]) if x.get('ean')}
        except (OSError, ValueError, KeyError):
            return {}

    @staticmethod
    def _cap3000():
        try:
            t = open(os.path.join(ROOT, 'crm', 'cap3000-data.js'), encoding='utf-8').read()
            return {m.group(2): m.group(1) for m in re.finditer(r'marque:"([^"]*)",ean:"(\d+)"', t)}
        except OSError:
            return {}

    def joli(self, m):
        """Graphie du catalogue quand la marque y existe déjà (« LA ROCHE POSAY » = « LA ROCHE-POSAY »)."""
        return self.canon.get(norm(m), m)

    def resoudre(self, c, des, candidats):
        """Rend (marque|'' , source). `candidats` = tous les « labos » lus pour ce produit, par priorité."""
        for l in candidats:
            if l and not est_fournisseur(l):
                return self.joli(l), 'autre colonne ou autre fichier du groupe'
        for nom, m in self.par_code.get(c, []):
            return self.joli(m), nom
        t = norm(des).split()
        for n in self.longueurs:
            if len(t) >= n and tuple(t[:n]) in self.table:
                return self.joli(self.table[tuple(t[:n])]), 'début de la désignation (table de marques connues)'
        return '', 'introuvable'


def court(x):
    """5.0 -> 5 et 0.0 -> 0 (quelques octets de moins par ligne ; la valeur ne change pas)."""
    return int(x) if float(x) == int(x) else x


# ───────────────────────── programme ─────────────────────────
def main():
    t0 = time.time()
    gcc = importer_regles()
    charger_pont()
    S = decouvrir()
    if 'POS' not in S['complet_ppht']:
        sys.exit('ARRÊT : aucun export complet avec ppht trouvé (POS attendu).')

    complet, cip7_map, complet_date = {}, {}, {}
    for site, d in list(S['complet_ppht'].items()) + [(s, d) for s, d in S['complet_sans_ppht'].items() if s not in S['complet_ppht']]:
        complet[site], c7 = lire_complet(d['f'], site in S['complet_ppht'], 'export ' + site)
        complet_date[site] = d['date']
        for k, v in c7.items():
            cip7_map.setdefault(k, v)
        log('  %s export complet%s : %d codes' % (site, '' if site in S['complet_ppht'] else ' (sans PPHT)', len(complet[site])))
    extraits = {s: lire_extrait(d['f'], s) for s, d in S['extraits'].items()}
    ext_date = {s: d['date'] for s, d in S['extraits'].items()}
    xls, xls_date = {}, None
    if S['xls']:
        xls = lire_xls(S['xls']['f'], cip7_map)
        xls_date = S['xls']['date']
    juin = lire_juin(S['juin']['f']) if S['juin'] else {}
    juin_date = S['juin']['date'] if S['juin'] else None
    log('  extraits %d sites · xls %d codes (%s) · juin %d codes' % (len(extraits), len(xls), xls_date, len(juin)))

    union = set()
    for d in list(complet.values()) + list(extraits.values()):
        union |= set(d)
    union |= set(xls)
    log('UNION :', len(union))

    repertoire = set(js_json(os.path.join(V2, 'generiqueurs-data.js')).keys())
    ps_txt = open(os.path.join(V2, 'prod-stats-data.js'), encoding='utf-8').read()
    ps = {}
    for o in json.loads(ps_txt[ps_txt.index('['):ps_txt.rindex(']') + 1]):
        c = canon(str(o['c']))
        if c not in ps or (o.get('n') or 0) > (ps[c].get('n') or 0):
            ps[c] = o
    connu_ray, ATC, saison, nr_cip, rayons_ref = charger_rayons()
    try:
        nature_afm = json.load(open(os.path.join(SRC, 'nature-afm.json'), encoding='utf-8'))
    except (OSError, ValueError):
        nature_afm = {}

    # ───────── fiche de chaque produit
    prods = {}
    stat_tarif = Counter()
    for c in sorted(union):
        pos, sep = complet.get('POS', {}).get(c), complet.get('SEP', {}).get(c)
        autres = [complet[s].get(c) for s in complet if s not in ('POS', 'SEP') and complet[s].get(c)]
        ex = {s: extraits[s][c] for s in extraits if c in extraits[s]}
        j, x = juin.get(c), xls.get(c)

        cand_d = [o['d'] for o in ([pos, sep] + autres) if o and o['d']]
        if x and x['d']:
            cand_d.append(x['d'])
        cand_d += [ex[s]['d'] for s in SITE_PRIO if s in ex and ex[s]['d']]
        if j:
            cand_d.append(j['d'])
        des = cand_d[0] if cand_d else ''
        if c in LIBELLES:
            des = LIBELLES[c]
        elif RP_RE.search(des):
            des = RP_RE.sub('', des).strip()
        des = re.sub(r'\s+', ' ', des).strip()
        des = re.sub(r'\s+ITP$', '', des)    # le sigle du dépositaire ne sort pas non plus, même en fin de désignation
        if not des:
            ANOM['produit sans désignation exploitable (écarté)'] += 1
            continue

        # tous les « labos » lus pour ce produit, par priorité (le premier = ce que les fichiers disaient jusqu'ici)
        cand_l = []
        for o in [pos] + ([j] if j else []) + [sep] + autres:
            if o:
                cand_l += [o['labo'], o.get('marque', '')]
        cand_l += [ex[s]['labo'] for s in SITE_PRIO if s in ex]
        cand_l = [l for l in cand_l if l]

        # tarif : PPHT de l'export POS, le même pour les 7 ; replis datés SEULEMENT si POS n'en a pas
        ppht, p_src = 0.0, None
        if pos and pos['ppht'] > 0:
            ppht, p_src = pos['ppht'], 'export POS (%s)' % complet_date['POS']
        else:
            for s in SITES:
                o = complet.get(s, {}).get(c)
                if o and o['ppht'] > 0 and s != 'POS':
                    ppht, p_src = o['ppht'], 'repli : export %s' % s
                    break
        if ppht <= 0 and j and j['ppht'] > 0:
            ppht, p_src = j['ppht'], 'repli : catalogue de juin (%s)' % juin_date
        if ppht <= 0 and ex:
            vals = defaultdict(list)
            for s in SITE_PRIO:
                if s in ex and ex[s]['ppht'] > 0:
                    vals[ex[s]['ppht']].append(s)
            if vals:
                best = max(vals.items(), key=lambda kv: (len(kv[1]), -SITE_PRIO.index(kv[1][0])))
                ppht, p_src = best[0], 'repli : extrait NR de site'
        stat_tarif[p_src.split(' (')[0] if p_src else 'AUCUN'] += 1

        afm = next((o['afm'] for o in [pos, sep] + autres if o and o['afm']), '')
        if not afm and ex:
            afm = next((ex[s]['afm'] for s in SITE_PRIO if s in ex and ex[s]['afm']), '')
        if not afm and j:
            afm = j['afm']
        if not afm and c in nature_afm:
            afm = nature_afm[c][0]
        # `artnature` : celle de la ligne POS quand le produit y est (même lecture que generate_catalogue_complet.py,
        # qui ne lit que cet export) ; sinon SEP, puis l'ancien catalogue.
        nat_brut = pos['nat'] if pos else (sep['nat'] if sep else (j['nat'] if j else ''))

        # état du stock par site : 2 = en stock · 1 = ligne au relevé du site, stock 0 · 0 = pas de relevé
        etat = {}
        for s in ('POS', 'SEP'):
            if s in complet:
                o = complet[s].get(c)
                etat[s] = '0' if o is None else ('2' if o['s'] > 0 else '1')
        for s in XLS_COLS.values():
            if x and xls_date:
                etat[s] = '2' if x['s'].get(s, 0) > 0 else '1'
            elif s in ex:
                # SOP : le fichier STOCK de septembre ne le liste pas -> pas en stock ; la ligne de l'extrait existe
                etat[s] = '1' if (xls_date and s == SITE_BASE_XLS) else ('2' if ex[s]['s'] > 0 else '1')
            else:
                etat[s] = '0'
        sites = ''.join(etat.get(s, '0') for s in SITES)

        prods[c] = {'c': c, 'd': des, 'cand': cand_l, 'mol': j['mol'] if j else '', 'afm': afm,
                    'remb': afm == 'REMBSS', 'nat': nat_brut, 'ppht': ppht, 'p_src': p_src, 'sites': sites,
                    'n': int(ps[c]['n']) if c in ps and ps[c].get('n') else 0}

    # ───────── MARQUE : jamais un fournisseur
    reels = [l for p in prods.values() for l in p['cand'][:1] if not est_fournisseur(l)]
    M = Marques(reels)
    avant_apres, src_marque = [], Counter()
    n_fourn = 0
    for c, p in prods.items():
        premier = p['cand'][0] if p['cand'] else ''
        avait = any(est_fournisseur(l) for l in p['cand'])
        marque, src = M.resoudre(c, p['d'], p['cand'])
        if avait:
            n_fourn += 1
            src_marque[src] += 1
            avant_apres.append((p['d'], premier, marque, src))
        elif not marque:
            src = 'aucun labo dans les fichiers'
        p['labo'] = marque
        p['labo_src'] = src if avait else None
        p['avait_fournisseur'] = avait
        assert not est_fournisseur(p['labo']), c

    # ───────── nature (fam) et net : règles de generate_catalogue_complet.py
    cat = {c: {'c': c, 'remb': p['remb'], 'labo': p['labo']} for c, p in prods.items()}
    labos_gen = gcc.labos_generiqueurs(cat, repertoire)
    chg, dangereux = Counter(), 0
    for c, p in prods.items():
        calc = gcc.famille(c, p['ppht'], p['remb'], p['nat'], p['d'], repertoire, p['labo'], labos_gen)
        f = calc
        # Comme generate_catalogue_complet.py : quand le réseau connaît le produit, SA famille est celle
        # que l'app affiche partout ailleurs, même si le statut du fichier récent a changé depuis.
        if c in ps and ps[c].get('f'):
            fr = ps[c]['f']
            f = fr
            if (fr == 'nr') == p['remb']:
                chg['statut réseau ≠ fichier récent (classement du réseau retenu)'] += 1
            if fr != calc:
                chg['famille du réseau retenue (différente du calcul)'] += 1
                if not gcc.porte_abandon(fr) and gcc.porte_abandon(calc):
                    dangereux += 1
        p['fam'] = f
        net = 0.0
        if c in ps and nombre(ps[c].get('net')) > 0:
            net = gcc.r2(ps[c]['net'])
        elif p['ppht'] > 0:
            net = gcc.r2(p['ppht'] - gcc.abandon_ip(p['ppht'])) if gcc.porte_abandon(f) else p['ppht']
        p['net'] = net

    # ───────── rayon
    rayons = list(rayons_ref)
    ray_idx = {l: i for i, l in enumerate(rayons)}
    src_ray = Counter()
    for c, p in prods.items():
        ray = None
        if c in connu_ray:
            ray, k = connu_ray[c], 'mkt-rayons-data.js'
        elif c in saison and saison[c].get('a') and saison[c]['a'][:3] in ATC:
            ray, k = ATC[saison[c]['a'][:3]], 'ATC Medic\'AM'
        elif c in nr_cip:
            ray, k = nr_cip[c], 'catégorie NR du mix'
        else:
            k = 'aucun'
        src_ray[k] += 1
        if ray is None:
            p['ray'] = -1
        else:
            if ray not in ray_idx:
                ray_idx[ray] = len(rayons)
                rayons.append(ray)
            p['ray'] = ray_idx[ray]

    # ───────── écriture
    ordre = sorted(prods.values(), key=lambda p: (-p['n'], p['d'], p['c']))
    labos, mols, rows = {}, {}, []
    for p in ordre:
        li = labos.setdefault(p['labo'], len(labos)) if p['labo'] else -1
        mi = mols.setdefault(p['mol'], len(mols)) if p['mol'] else -1
        rows.append([p['c'], p['d'], li, mi, p['fam'], court(p['ppht']), court(p['net']), p['n'], p['ray'], p['sites']])
    labos_l = [k for k, _ in sorted(labos.items(), key=lambda kv: kv[1])]
    mols_l = [k for k, _ in sorted(mols.items(), key=lambda kv: kv[1])]

    genere = time.strftime('%Y-%m-%d')
    fraich = {}
    for s in SITES:
        if s in complet:
            fraich[s] = {'date': complet_date[s], 'quoi': 'stock (export complet)'}
        elif xls_date:
            q = 'stock (fichier STOCK)'
            if s in ext_date:
                q += ' ; repli extrait non remboursables du ' + ext_date[s]
            fraich[s] = {'date': xls_date, 'quoi': q}
        elif s in ext_date:
            fraich[s] = {'date': ext_date[s], 'quoi': 'stock (extrait non remboursables)'}
    meta = {
        'n': len(rows), 'genere': genere,
        'tarif': {'source': 'export POS', 'date': complet_date['POS'], 'note': 'PPHT commun aux 7 établissements'},
        'sites': SITES, 'fraicheur': fraich,
        'fournisseurs': sorted(FOURNISSEURS_VUS),
        'cols': ['cip', 'd', 'labo', 'mol', 'fam', 'ppht', 'net', 'n', 'rayon', 'sites'],
    }
    for k in FORBIDDEN:
        assert k not in json.dumps(meta, ensure_ascii=False).lower(), 'clé interdite dans meta : ' + k
    dump = lambda o: json.dumps(o, ensure_ascii=False, separators=(',', ':'))
    js = ('/* Catalogue du groupe (7 établissements) — généré par generate_catalogue_groupe.py. '
          'Ne pas éditer à la main. %d références au %s. */\n'
          'var CATALOGUE_GROUPE = {\nmeta: %s,\nlabos: %s,\nmols: %s,\nrayons: %s,\nrows: [\n%s\n]};\n'
          % (len(rows), genere, dump(meta), dump(labos_l), dump(mols_l), dump(rayons),
             ',\n'.join(dump(r) for r in rows)))
    open(OUT, 'w', encoding='utf-8').write(js)

    # ───────── mesures (reprises dans le rapport)
    import gzip
    brut = os.path.getsize(OUT)
    gz = len(gzip.compress(js.encode('utf-8'), 9))
    nrb = sum(1 for p in prods.values() if p['remb'])
    log('\n=== MESURES ===')
    log('n total %d · remboursables %d · NR %d' % (len(rows), nrb, len(rows) - nrb))
    log('poids brut %.2f Mo (%d o) · gzip %.2f Mo (%d o)' % (brut / 1048576, brut, gz / 1048576, gz))
    log('en stock par site : ' + ' · '.join('%s %d' % (s, sum(1 for r in rows if r[9][i] == '2')) for i, s in enumerate(SITES)))
    log('relevé absent (0) par site : ' + ' · '.join('%s %d' % (s, sum(1 for r in rows if r[9][i] == '0')) for i, s in enumerate(SITES)))
    log('sans tarif %d · sans rayon %d · sans marque %d · sans molécule %d' % (
        sum(1 for r in rows if r[5] <= 0), sum(1 for r in rows if r[8] < 0), sum(1 for r in rows if r[2] < 0), sum(1 for r in rows if r[3] < 0)))
    log('tarif par source : ' + ' · '.join('%s %d' % kv for kv in stat_tarif.most_common()))
    log('rayon par source : ' + ' · '.join('%s %d' % kv for kv in src_ray.most_common()))
    log('familles : ' + ' · '.join('%s %d' % kv for kv in Counter(r[4] for r in rows).most_common()))
    log('changements de classement : ' + '; '.join('%s : %d' % kv for kv in chg.items()) + ' ; dangereux (réseau sans abandon, calcul avec) : %d' % dangereux)
    log('graphies de fournisseur lues : %s' % sorted(FOURNISSEURS_VUS))
    log('produits qui avaient un fournisseur comme « labo » : %d' % n_fourn)
    for k, v in src_marque.most_common():
        log('   marque retrouvée par : %-60s %d' % (k, v))
    log('--- sans marque retrouvée (désignations)')
    for d, av, ap, src in avant_apres:
        if not ap:
            log('   ' + d[:70])
    log('--- 20 exemples avant -> après')
    pas = max(1, len(avant_apres) // 20)
    for d, av, ap, src in avant_apres[::pas][:20]:
        log('   %-48s | %-18s -> %-22s | %s' % (d[:48], av, ap or '(rien)', src))
    log('anomalies : ' + '; '.join('%s : %d' % kv for kv in sorted(ANOM.items())))
    log('terminé en %.0f s' % (time.time() - t0))


if __name__ == '__main__':
    main()
