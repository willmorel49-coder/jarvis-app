#!/usr/bin/env python3
"""Base clients (export « clients actifs » du système de gestion) → crm/v2/clients-actifs.js

Source : STATS/total ventes/clients_actifs_<date>.xlsx (le plus récent), 2 077 tiers
actifs : téléphone, portable, e-mail, contact (titulaire), fonction, logiciel de
gestion (LGPI / Winpharma / Offilog / Alliadis / Pharmaland / autre), enseigne
(= groupement déclaré dans la base clients, colonne TIRENSEIGNE), commercial
(REPCODE), UGA, rythme de livraison (TYPO).

⚠️ Fichier PROTÉGÉ (.gitignore) : c'est NOTRE liste de clients avec le commercial
qui suit chacun — même traitement que wml-officines-ca.js. L'encours (TRQENCOURS)
n'est pas repris du tout. Dépôt : seau Supabase donnees-protegees, sous le nom
clients-actifs.js (voir skill donnees-protegees-jarvis), puis bump du jeton V.

Clé = TIRCODE = code officine des fichiers de ventes (WML_OFFICINES.id).
Seuls les tiers « Pharmacie » sont repris : distributeurs, labos et revendeurs
n'ont ni fiche officine ni groupement.

11/09/2026 — s'y ajoute le fichier clients d'ESCALE PHARMA (Chilly-Mazarin,
commerciaux GUY / TIF / PHI / GER) : STATS/total ventes/liste_clients_escale_<date>.xlsx,
colonnes en minuscules (tircode, adrtel, adrmail, adrcontactnom/prenom, lgo,
tircible1 = groupement, repcode, generiqueur1/2). ⚠️ tirenseigne y vaut
« OFFRE INTEGRAL » pour tous : ce n'est PAS le groupement, on lit tircible1.
Quand tircible1 est vide, la base groupements (~/JARVIS/GROUPEMENTS, scrapée
2026) comble par CP + nom exact, uniquement un groupement_actuel de statut
« changé » / « confirmé » — jamais « non_retrouvé » (= PHIRST 2025, périmé).
Un 11ᵉ champ (génériqueur·s) est ajouté à chaque ligne ; vide côté Intégral.

Usage : /usr/bin/python3 generate_clients_actifs.py
"""
import glob, io, json, os, re, sys, datetime, sqlite3, unicodedata
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
STATS = '/Users/williammorel/JARVIS/APP/STATS/total ventes'
OUT = os.path.join(BASE, 'crm/v2/clients-actifs.js')
LGO = [('LGPI', 'LGPI'), ('WINPHARMA', 'Winpharma'), ('OFFILOG', 'Offilog'),
       ('ALLIADIS', 'Alliadis'), ('PHARMALAND', 'Pharmaland'), ('AUTRES', 'Autre')]
# Libellés de la colonne `lgo` du fichier Escale → même orthographe que ci-dessus
LGO_ESCALE = {'LGPI': 'LGPI', 'WINPHARMA': 'Winpharma', 'SMART RX': 'Smart Rx', 'LEO': 'LEO',
              'PHARMALAND': 'Pharmaland', 'LSI PHARMALAND': 'Pharmaland', 'ALLIADIS': 'Alliadis',
              'ACTIPHARM': 'Actipharm', 'VINDALIS': 'Vindilis', 'VINPILIS': 'Vindilis'}
GROUPEMENTS_DB = '/Users/williammorel/JARVIS/GROUPEMENTS/data/output/pharmacies.sqlite'


def norm_nom(v):
    return re.sub(r'[^A-Z0-9]', '', unicodedata.normalize('NFD', s(v)).encode('ascii', 'ignore').decode().upper())


def groupements_db():
    """(cp, nom normalisé) → groupement actuel, seulement quand il est sûr et unique."""
    idx = {}
    try:
        db = sqlite3.connect(GROUPEMENTS_DB)
        for nom, cp, g in db.execute("select nom, cp, groupement_actuel from pharmacies "
                                     "where groupement_actuel != '' and statut_groupement in ('changé', 'confirmé')"):
            k = (s(cp), norm_nom(nom))
            idx.setdefault(k, set()).add(s(g))
    except Exception as e:
        print('⚠️ base groupements illisible (%s) : aucun comblement' % e)
        return {}
    return {k: next(iter(v)) for k, v in idx.items() if len(v) == 1}


def lire_escale(d):
    """Fichier clients Escale Pharma → mêmes lignes que l'export Intégral, fusionnées dans d."""
    fichiers = sorted(glob.glob(os.path.join(STATS, 'liste_clients_escale_*.xlsx')))
    if not fichiers:
        print('⚠️ aucun liste_clients_escale_*.xlsx : base Escale non reprise')
        return None
    src = fichiers[-1]
    ws = openpyxl.load_workbook(src, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    H = {s(h).lower(): i for i, h in enumerate(next(rows))}
    for c in ('tircode', 'tirsociete', 'adrcodepostal', 'adrtel', 'adrmail', 'adrcontactnom', 'tircible1', 'repcode', 'lgo'):
        if c not in H:
            sys.exit('fichier Escale : colonne manquante : ' + c)
    gdb = groupements_db()
    n, n_grp, n_comble, n_new, n_fusion = 0, 0, 0, 0, 0
    for r in rows:
        code = s(r[H['tircode']])
        if not code or s(r[H.get('tiractivite', H['tircode'])]) not in ('Pharmacie', code):
            continue
        n += 1
        grp = s(r[H['tircible1']])
        if grp:
            n_grp += 1
        else:
            grp = gdb.get((s(r[H['adrcodepostal']]), norm_nom(r[H['tirsociete']])), '')
            if grp:
                n_comble += 1
        nom, pre = s(r[H['adrcontactnom']]), s(r[H.get('adrcontactprenom', H['adrcontactnom'])])
        gen = [s(r[H[c]]).title() for c in ('generiqueur1', 'generiqueur2') if c in H and s(r[H[c]])]
        ligne = [
            tel(r[H['adrtel']]), tel(r[H.get('adrportable', H['adrtel'])]), s(r[H['adrmail']]).lower(),
            ' '.join(x for x in (pre.title(), nom.upper()) if x),
            '', LGO_ESCALE.get(s(r[H['lgo']]).upper(), s(r[H['lgo']]).title()),
            grp, s(r[H['repcode']]), '', '', ' / '.join(gen),
        ]
        if code not in d:
            d[code] = ligne
            n_new += 1
        else:   # déjà connue par l'export Intégral : on ne comble que les vides
            n_fusion += 1
            cur = d[code]
            for i, v in enumerate(ligne):
                if not cur[i] and v:
                    cur[i] = v
    print('Escale   : %s — %d officines (groupement déclaré %d, comblé par la base groupements %d, sans %d) ; %d nouvelles, %d déjà connues'
          % (os.path.basename(src), n, n_grp, n_comble, n - n_grp - n_comble, n_new, n_fusion))
    return src


def s(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v)).strip()


def tel(v):
    """'494012345' (Excel a mangé le 0) → '04 94 01 23 45' ; les textes sont gardés."""
    t = s(v)
    if not t:
        return ''
    d = re.sub(r'\D', '', t)
    if isinstance(v, (int, float)) and len(d) == 9:
        d = '0' + d
    if len(d) == 10:
        return ' '.join(d[i:i + 2] for i in range(0, 10, 2))
    return t


def contact(r, H):
    civ, nom, pre = s(r[H['ADRCONTACTTYPE']]), s(r[H['ADRCONTACTNOM']]), s(r[H['ADRCONTACTPRENOM']])
    if not nom and not pre:
        return ''
    return ' '.join(x for x in (civ, pre.title() if pre else '', nom.upper()) if x)


def fonction(v):
    f = s(v).lower()
    if not f:
        return ''
    if f.startswith('titulaire'):
        return 'Titulaire'
    return s(v)[:1].upper() + s(v)[1:]


def main():
    fichiers = sorted(glob.glob(os.path.join(STATS, 'clients_actifs_*.xlsx')))
    if not fichiers:
        sys.exit('aucun clients_actifs_*.xlsx dans ' + STATS)
    src = fichiers[-1]
    date = re.search(r'(\d{4}-\d{2}-\d{2})', os.path.basename(src))
    date = date.group(1) if date else str(datetime.date.today())
    ws = openpyxl.load_workbook(src, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    H = {h: i for i, h in enumerate(next(rows))}
    for c in ('TIRCODE', 'TIRACTIVITE', 'ADRTEL', 'ADRMAIL', 'ADRCONTACTNOM', 'TIRENSEIGNE', 'REPCODE', 'LGPI'):
        if c not in H:
            sys.exit('colonne manquante : ' + c)
    # Une même pharmacie apparaît jusqu'à 3 fois : une ligne par société du
    # groupe qui la livre (colonne Structure). On FUSIONNE par code : première
    # valeur remplie pour chaque champ, logiciels réunis, commerciaux distincts.
    d, n_pharma, n_autres = {}, 0, 0
    for r in rows:
        code = s(r[H['TIRCODE']])
        if not code:
            continue
        if s(r[H['TIRACTIVITE']]) != 'Pharmacie':
            n_autres += 1
            continue
        n_pharma += 1
        lgo = [lib for col, lib in LGO if col in H and s(r[H[col]]) == 'O']
        ligne = [
            tel(r[H['ADRTEL']]),                 # 0 téléphone
            tel(r[H['ADRPORTABLE']]),            # 1 portable
            s(r[H['ADRMAIL']]).lower(),          # 2 e-mail
            contact(r, H),                       # 3 contact (civilité Prénom NOM)
            fonction(r[H['ADRCONTACTFONCTION']]),  # 4 fonction
            lgo,                                 # 5 logiciel(s) de gestion
            s(r[H['TIRENSEIGNE']]),              # 6 enseigne = groupement (base clients)
            [s(r[H['REPCODE']])] if s(r[H['REPCODE']]) else [],  # 7 commercial (code)
            s(r[H['UGA']]),                      # 8 UGA
            s(r[H['TYPO']]).capitalize(),        # 9 rythme de livraison
            '',                                  # 10 génériqueur(s) — fichier Escale seulement
        ]
        if code not in d:
            d[code] = ligne
            continue
        cur = d[code]
        for i, v in enumerate(ligne):
            if i in (5, 7):
                cur[i] = cur[i] + [x for x in v if x not in cur[i]]
            elif not cur[i] and v:
                cur[i] = v
    for v in d.values():
        v[5] = ' / '.join(v[5])
        v[7] = ' / '.join(v[7])
    src_escale = lire_escale(d)
    txt = ('// Intégral Pharma — base clients (export clients actifs du %s) — généré le %s\n'
           '// ⚠️ NE JAMAIS COMMITER. Servi par adresse signée (Supabase).\n'
           '// code officine → [tel, portable, email, contact, fonction, logiciel, enseigne, commercial, uga, livraison, generiqueurs]\n'
           'window.CLIENTS_ACTIFS = {date:%s, n:%d, d:%s};\n'
           % (date, datetime.date.today(), json.dumps(date), len(d), json.dumps(d, ensure_ascii=False)))
    io.open(OUT, 'w', encoding='utf-8').write(txt)
    # se relire : le fichier écrit est bien celui annoncé
    relu = io.open(OUT, encoding='utf-8').read()
    m = re.search(r'window\.CLIENTS_ACTIFS = (\{.*\});\s*$', relu, re.S)
    obj = json.loads(re.sub(r'^\{date:', '{"date":', re.sub(r', n:', ', "n":', re.sub(r', d:', ', "d":', m.group(1)))))
    assert obj['n'] == len(obj['d']) == len(d), 'relecture ≠ écriture'
    print('source   : %s%s' % (os.path.basename(src), (' + ' + os.path.basename(src_escale)) if src_escale else ''))
    print('lignes pharmacie : %d → %d officines distinctes (autres tiers ignorés : %d)' % (n_pharma, len(d), n_autres))
    print('avec tel %d · mail %d · contact %d · logiciel %d · enseigne %d' % tuple(
        sum(1 for v in d.values() if v[i]) for i in (0, 2, 3, 5, 6)))
    print('écrit    : %s (%d Ko)' % (OUT, os.path.getsize(OUT) // 1024))


if __name__ == '__main__':
    main()
