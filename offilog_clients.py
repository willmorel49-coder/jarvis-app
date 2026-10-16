#!/usr/bin/env python3
"""Pharmacies clientes d'Offilog — colonne OFFILOG des fichiers de géolocalisation (13/10/2026).

Remontée de Karine du 24/09/2026 : « infos clients : sont-ils Offilog ? ».
La réponse est dans STATS/<PRÉFIXE>_geolocalisation_<date>.xlsx, colonne OFFILOG (« O » / « N »),
clé TIRCODE. La colonne se lit PAR SON NOM d'en-tête, jamais par son rang.

Règles (arrêtées avec Will le 13/10/2026) :
  - dernier fichier de chaque commercial (préfixe) seulement ;
  - case vide = on ne sait pas : la pharmacie est ABSENTE de la table, jamais « non » par défaut ;
  - deux lignes qui se contredisent pour une même pharmacie : « O » l'emporte ;
  - la date annoncée à l'écran est celle du NOM du fichier, jamais écrite en dur.

C'est une liste de clients : elle ne va JAMAIS dans un fichier public. Ce script écrit
`wml-offilog.json` (couvert par .gitignore) ; `proteger_conditions.py` le range dans la
table protégée `wml-officines-ca.js` (Supabase), sous la clé `off`.

16/10/2026 — l'offre Sanofi / UPSA suit le même chemin (demande de Will : « ajoute si les clients ont
sanofi upsa en offre ou offilog »). Source : colonne TIRBRANCHE (lue PAR SON NOM), clé TIRCODE ;
SANOFI23 = Sanofi, UPSA23 = UPSA, SANOFI UPSA25 = les deux, vide ou autre valeur (APM, SRA, G1000…) =
ni l'un ni l'autre. Fichiers lus : le dernier STATS/<PRÉFIXE>_geolocalisation_<date>.xlsx de chaque
préfixe ET le dernier STATS/total ventes/clients_actifs_<date>.xlsx. Un fichier SANS colonne TIRBRANCHE
(les 4 Escale du 11/09) laisse ses pharmacies ABSENTES. Par pharmacie le fichier le plus récent l'emporte ;
à date égale, UNION des lignes (Sanofi oui sur une ligne = Sanofi oui). Une case vide dans un fichier qui
A la colonne = « pas en offre ». Forme : dans l'objet `off` existant, deux clés en plus de d et m :
bd = [dates des fichiers], b = {TIRCODE: [code, rang de la date dans bd]}, code 0 aucun / 1 Sanofi /
2 UPSA / 3 les deux. `d` et `m` (Offilog) ne changent pas.

Usage : python3 offilog_clients.py [dossier STATS]   (appelé aussi par generate_wml_v2.py)
"""
import glob, io, json, os, re, sys

BASE = os.path.dirname(os.path.abspath(__file__))
SORTIE = os.path.join(BASE, 'wml-offilog.json')


def _cle(code):
    try:
        return str(int(float(code)))
    except (TypeError, ValueError):
        m = re.match(r'\d+', str(code or ''))
        return m.group(0) if m else str(code or '').strip()


def lire(stats):
    """-> {'d': [dates des fichiers], 'm': {TIRCODE: [1 si « O » / 0 si « N », rang de la date]}}"""
    import openpyxl
    derniers = {}
    for p in glob.glob(os.path.join(stats, '*_geolocalisation_*.xlsx')):
        pref, date = os.path.basename(p)[:-5].split('_geolocalisation_')
        if pref not in derniers or date > derniers[pref][0]:
            derniers[pref] = (date, p)
    vus = {}   # TIRCODE -> (valeur, date)
    for date, p in sorted(derniers.values()):
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True); ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        h = [str(c).strip() if c is not None else '' for c in next(it)]
        if 'TIRCODE' not in h or 'OFFILOG' not in h:
            print('  [offilog] pas de colonne OFFILOG :', os.path.basename(p))
            wb.close(); continue
        ti, oi = h.index('TIRCODE'), h.index('OFFILOG')
        for r in it:
            if not r or ti >= len(r) or not r[ti]:
                continue
            v = str(r[oi]).strip().upper() if oi < len(r) and r[oi] is not None else ''
            if v not in ('O', 'N'):
                continue
            k = _cle(r[ti]); avant = vus.get(k)
            # « O » l'emporte sur « N » ; à valeur égale, le fichier le plus récent donne la date
            if avant is None or (v == 'O' and avant[0] == 'N') or (v == avant[0] and date > avant[1]):
                vus[k] = (v, date)
        wb.close()
    dates = sorted({d for _, d in vus.values()})
    return {'d': dates, 'm': {k: [1 if v == 'O' else 0, dates.index(d)] for k, (v, d) in vus.items()}}


def _code_branche(v):
    """TIRBRANCHE -> 0 aucun / 1 Sanofi / 2 UPSA / 3 les deux (sous-chaînes, majuscules, espaces rognés)."""
    v = str(v).strip().upper() if v is not None else ''
    return (1 if 'SANOFI' in v else 0) | (2 if 'UPSA' in v else 0)


def lire_offre(stats):
    """-> {'bd': [dates des fichiers], 'b': {TIRCODE: [code, rang de la date]}} ; {} si rien de lu."""
    import openpyxl
    derniers = {}
    for p in glob.glob(os.path.join(stats, '*_geolocalisation_*.xlsx')):
        pref, date = os.path.basename(p)[:-5].split('_geolocalisation_')
        if pref not in derniers or date > derniers[pref][0]:
            derniers[pref] = (date, p)
    fichiers = list(derniers.values())
    ca = []
    for p in glob.glob(os.path.join(stats, 'total ventes', 'clients_actifs_*.xlsx')):
        ca.append((os.path.basename(p)[:-5].split('clients_actifs_')[1], p))
    if ca:
        fichiers.append(max(ca))
    vus = {}   # TIRCODE -> (code, date)
    for date, p in sorted(fichiers):
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True); ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        h = [str(c).strip() if c is not None else '' for c in next(it)]
        if 'TIRCODE' not in h or 'TIRBRANCHE' not in h:
            print('  [offre] pas de colonne TIRBRANCHE :', os.path.basename(p))
            wb.close(); continue
        ti, bi = h.index('TIRCODE'), h.index('TIRBRANCHE')
        dans = {}   # lignes de CE fichier : union par pharmacie
        for r in it:
            if not r or ti >= len(r) or not r[ti]:
                continue
            k = _cle(r[ti])
            dans[k] = dans.get(k, 0) | _code_branche(r[bi] if bi < len(r) else None)
        wb.close()
        for k, c in dans.items():
            avant = vus.get(k)
            if avant is None or date > avant[1]:
                vus[k] = (c, date)
            elif date == avant[1]:
                vus[k] = (avant[0] | c, date)
    dates = sorted({d for _, d in vus.values()})
    if not vus:
        return {}
    return {'bd': dates, 'b': {k: [c, dates.index(d)] for k, (c, d) in vus.items()}}


def ecrire(stats):
    t = lire(stats)
    if not t['m']:
        # rien de lu (dossier absent, colonne disparue) : on garde la table d'avant, on le dit
        print('  [offilog] 🔴 aucune valeur lue dans %s : wml-offilog.json laissé tel quel' % stats)
        return t
    o = lire_offre(stats)
    if not o:
        # l'offre ne lit rien : on garde celle de la table d'avant si elle en avait une, et on le dit
        print('  [offre] 🔴 aucune valeur TIRBRANCHE lue dans %s : offre Sanofi / UPSA conservée telle quelle' % stats)
        try:
            avant = json.loads(io.open(SORTIE, encoding='utf-8').read())
            o = {k: avant[k] for k in ('bd', 'b') if k in avant}
        except (IOError, OSError, ValueError):
            o = {}
    t.update(o)
    io.open(SORTIE, 'w', encoding='utf-8').write(json.dumps(t, separators=(',', ':'), sort_keys=True))
    oui = sum(1 for v in t['m'].values() if v[0])
    print('  [offilog] %d pharmacies : %d « O », %d « N » (fichiers du %s)'
          % (len(t['m']), oui, len(t['m']) - oui, ', '.join(t['d'])))
    if t.get('b'):
        n = [sum(1 for v in t['b'].values() if v[0] == c) for c in range(4)]
        print('  [offre] %d pharmacies : %d Sanofi seul, %d UPSA seul, %d les deux, %d aucune (fichiers du %s)'
              % (len(t['b']), n[1], n[2], n[3], n[0], ', '.join(t['bd'])))
    return t


if __name__ == '__main__':
    ecrire(sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, 'STATS'))
