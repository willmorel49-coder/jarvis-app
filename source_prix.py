# -*- coding: utf-8 -*-
"""Choix COMMUN de la source de prix — une seule logique pour tous les générateurs (14/10/2026).

Règle (posée dans generate_ppht.py le 28/09/2026) : la source se choisit sur ce qu'elle
CONTIENT, jamais sur son nom ni sur une date lue dans son nom. On prend le fichier
STATS/*.xlsx le PLUS RÉCENT qui porte les colonnes `artcodebarre` ET `ppht`.
(`stock POS.xlsx` du 21/09 la porte ; `stock SEP 17092026.xlsx` n'a que `atfprix` : écarté.)

Les quatre générateurs qui choisissaient encore par le nom (prod_stats, catalogue_complet,
marketing_mix, mol_stats) lisaient le tarif du 22/06 : 2 001 prix différents sur 15 231
produits communs (Slinda 26,67 au lieu de 26,16).

Colonne absente de la source récente (`artmol`, `artcategorie` : `stock POS.xlsx` ne les a
pas) : `colonne_repli()` la reprend, PAR CODE-BARRES, dans le fichier le plus récent qui la
porte (aujourd'hui `stock et prix 22 06 2026.xlsx`). Un produit né après juin n'y figure pas :
sa valeur reste vide plutôt qu'inventée. Python 3.9.
"""
import glob
import os

import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))


def _colonnes(chemin):
    try:
        w = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
        h = [c.value for c in next(w.active.iter_rows(min_row=1, max_row=1))]
        w.close()
        return set(x for x in h if x)
    except Exception:
        return set()


def _candidats():
    c = [f for f in glob.glob(os.path.join(ROOT, 'STATS', '*.xlsx'))
         if not os.path.basename(f).startswith('~$')]
    c.sort(key=os.path.getmtime, reverse=True)
    return c


def source_prix(requises=('artcodebarre', 'ppht'), verbose=True):
    """Chemin du fichier STATS/*.xlsx le plus récent qui porte toutes les colonnes `requises`."""
    for f in _candidats():
        if set(requises) <= _colonnes(f):
            if verbose:
                print('  [source prix] %s' % os.path.relpath(f, ROOT))
            return f
    raise SystemExit('ARRÊT : aucun fichier de STATS/ ne porte les colonnes %s.' % ' + '.join(requises))


def colonne_repli(nom, exclure=None):
    """{code-barres: valeur} de la colonne `nom`, lue dans le fichier le plus récent (autre que
    `exclure`) qui porte `artcodebarre` ET `nom`. Rend {} (avec un message) si aucun."""
    for f in _candidats():
        if exclure and os.path.abspath(f) == os.path.abspath(exclure):
            continue
        if {'artcodebarre', nom} <= _colonnes(f):
            ws = openpyxl.load_workbook(f, read_only=True, data_only=True).active
            h = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            ci, vi = h.index('artcodebarre'), h.index(nom)
            out = {}
            for r in ws.iter_rows(min_row=2, values_only=True):
                c = str(r[ci] or '').strip()
                v = r[vi]
                if c.isdigit() and len(c) >= 12 and v not in (None, '', '#N/A'):
                    out[c] = v
            print('  [repli %s] %d valeurs reprises de %s' % (nom, len(out), os.path.relpath(f, ROOT)))
            return out
    print('  [repli %s] aucun fichier ne porte cette colonne' % nom)
    return {}


def ppht_reference():
    """{CIP13: PPHT} de crm/v2/ppht-data.js (window.PPHT, surcharge STOCK 150926 comprise) : LA référence
    de prix de l'appli (décision du 14/10/2026 : un produit = un prix sur tous les écrans). {} si absent."""
    import json
    import re
    p = os.path.join(ROOT, 'crm', 'v2', 'ppht-data.js')
    try:
        t = open(p, encoding='utf-8').read()
        return json.loads(re.search(r'window\.PPHT\s*=\s*(\{.*?\});\n', t, re.S).group(1))
    except (OSError, AttributeError, ValueError):
        print('  ⚠️ ppht-data.js illisible : prix pris dans la source brute')
        return {}
