# -*- coding: utf-8 -*-
"""Ventes par ÉTABLISSEMENT et par produit (France / export) depuis janvier -> crm/v2/etab-ventes-data.js

Demande de Will (23/10/2026) : « on saura qui a quoi comme stock en gros et ce qui part à l'export
par établissement ». Alimente le bloc « Stock par établissement » du panneau produit (v2-comptoir.js).

Lit (lecture seule) ~/JARVIS/APP/STATS/<CODE>_<MM>_2026.xlsx et STATS/hors-perimetre/ (jamais les
*_geolocalisation_*, *_agregation*, ni les zips). Établissement = colonne __SOURCE_FILE__ ramenée à son
code (OPS_Pharma.xlsx -> OPS). Escale (GER, GUY, PHI, TIF) n'a pas de __SOURCE_FILE__ : ce n'est pas un
des 7 établissements, il est compté à part dans le journal et JAMAIS inventé dans le fichier.

RÈGLE D'EXPORT (analyse du 23/10/2026, ~/jarvis-preuves/ventes-etablissements-2026-10-23/RAPPORT.md §3) :
  une ligne est export si son client est classé « Revendeur » ou « Négociant » dans le plus récent
  STATS/total ventes/clients_actifs_*.xlsx (colonne TIRACTIVITE), OU si la ligne a TVATAUX = 0 et que son
  client n'est pas classé « Pharmacie ».

Sortie : window.ETAB_VENTES = { de, a, genere, regleExport, regleDepuis, lignes, fichiers,
  sites: {SITE: {depuis: 'AAAA-MM'}}, v: { cip13: { SITE: [unitesFrance, unitesExport, 'AAAA-MM-JJ'] } } }
  - a = date de la dernière facture LUE (pas la date du jour) ; de = 1er jour du premier mois lu.
  - depuis (par site) = premier mois où le site a un volume normal : au moins 20 % de sa médiane mensuelle
    de lignes (médiane des mois où il a des lignes). Les totaux d'un site ne comptent QUE depuis ce mois :
    le chiffre affiché et la période annoncée sont les mêmes.
  - unités = PLVQTE nette des avoirs. Dernière vente = dernière ligne à quantité > 0 hors avoir (AC...).
  - clé = CIP13 après le pont de codes (pont-codes.json, le même que le stock : un produit = un code).

⚠️ DONNÉE PROTÉGÉE (chiffres par établissement) : le fichier est dans .gitignore, jamais suivi par git ;
il se dépose À LA MAIN dans l'espace Supabase fermé (voir la skill donnees-protegees-jarvis), nom
`etab-ventes-data.js` à la racine du seau, et se charge à la demande (V2.loadFiles(['etabventes'])).

Le générateur REFUSE d'écrire si : un mois entre janvier et le dernier mois lu n'a aucun fichier ; un site
n'a aucune ligne un mois après son `depuis` ; il lit moins de 80 % des lignes de la fois précédente ; la
dernière facture lue est plus ancienne que celle de la fois précédente.

100 % local. /usr/bin/python3 (3.9) + openpyxl.   Usage : /usr/bin/python3 generate_etab_ventes.py
"""
import datetime
import glob
import json
import os
import re
import statistics
import sys
import time
from multiprocessing import Pool

import openpyxl

STATS = os.path.expanduser('~/JARVIS/APP/STATS')
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'crm', 'v2', 'etab-ventes-data.js')

SITES = ['CPR', 'HP', 'MSP', 'OPS', 'POS', 'SEP', 'SOP']      # les 7 de ETABS (v2-comptoir.js)
ESCALE = ['GER', 'GUY', 'PHI', 'TIF']                          # fichiers sans __SOURCE_FILE__
EXPORT_ACTIVITES = ('Revendeur', 'Négociant')
SEUIL_DEPUIS = 0.20        # part de la médiane mensuelle de lignes
SEUIL_RELECTURE = 0.80     # part des lignes de la fois précédente
REGLE_EXPORT = ('client Revendeur ou Négociant (clients_actifs) OU ligne à TVATAUX = 0 d\'un client '
                'non classé Pharmacie')
REGLE_DEPUIS = 'premier mois à au moins 20 % de la médiane mensuelle de lignes du site'
RX_FICHIER = re.compile(r'^([A-Z_]+?)_(\d\d)_2026\.xlsx$')
COLS = ['PCVNUM', 'PCVDATEEFFET', 'TIRCODE', 'PLVQTE', 'TVATAUX', 'ARTCODEBARRE', '__SOURCE_FILE__']


# ── Règles pures (testées dans tests/etab-ventes.test.py) ───────────────────
def code_tir(v):
    try:
        return str(int(float(v)))
    except (TypeError, ValueError):
        return '' if v is None else str(v).strip()


def code_cip(v):
    if v is None or v != v:
        return ''
    try:
        return str(int(float(v)))
    except (TypeError, ValueError):
        return str(v).strip()


def code_site(src):
    """OPS_Pharma / « OPS juillet_Pharma » -> OPS. Vide -> None (fichiers Escale)."""
    if src is None or src != src or not str(src).strip():
        return None
    return re.split(r'[ _]', str(src).strip())[0].upper()


def est_export(activite, tva):
    """Règle d'export. activite = TIRACTIVITE du client (None si client absent de clients_actifs)."""
    if activite in EXPORT_ACTIVITES:
        return True
    try:
        t0 = tva is not None and float(tva) == 0
    except (TypeError, ValueError):
        t0 = False
    return bool(t0 and activite != 'Pharmacie')


def jour(v, repli):
    """Date de la ligne en 'AAAA-MM-JJ' ; sinon le 1er du mois du fichier (repli)."""
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, str) and re.match(r'^\d{4}-\d\d-\d\d', v):
        return v[:10]
    return repli


def premier_mois_normal(lignes_par_mois):
    """lignes_par_mois : {mois(int): nb de lignes}. -> premier mois >= 20 % de la médiane (mois non vides)."""
    vals = [n for n in lignes_par_mois.values() if n > 0]
    if not vals:
        return None
    seuil = SEUIL_DEPUIS * statistics.median(vals)
    for m in sorted(lignes_par_mois):
        if lignes_par_mois[m] >= seuil and lignes_par_mois[m] > 0:
            return m
    return None


# ── Lecture d'un fichier (tourne dans un processus enfant) ──────────────────
def lire_fichier(t):
    """t = (chemin, code_fichier, mois, activites{tir: activité}, alias{cip: cip gardé}).
    -> agrégat compact du fichier, jamais les lignes."""
    chemin, fcode, mois, activites, alias = t
    repli = '2026-%02d-01' % mois
    res = {'fcode': fcode, 'mois': mois, 'lignes': 0, 'erreur': None, 'sites': {}, 'cips': {},
           'escale': 0, 'inconnu': {}, 'sans_cip': 0, 'dmax': None, 'dmin': None}
    try:
        wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
        it = wb.active.iter_rows(values_only=True)
        h = [str(x).strip() if x is not None else '' for x in next(it)]
        hi = {k: i for i, k in enumerate(h)}
        idx = [hi.get(c) for c in COLS]
        for r in it:
            if all(v is None for v in r):
                continue
            ligne = {c: (r[i] if i is not None and i < len(r) else None) for c, i in zip(COLS, idx)}
            res['lignes'] += 1
            site = code_site(ligne['__SOURCE_FILE__'])
            if site is None:
                site = 'ESCALE' if fcode in ESCALE else '?'
            if site == 'ESCALE':
                res['escale'] += 1
                continue
            if site not in SITES:
                res['inconnu'][site] = res['inconnu'].get(site, 0) + 1
                continue
            res['sites'][site] = res['sites'].get(site, 0) + 1
            try:
                q = float(ligne['PLVQTE'] or 0)
            except (TypeError, ValueError):
                q = 0.0
            d = jour(ligne['PCVDATEEFFET'], repli)
            if res['dmax'] is None or d > res['dmax']:
                res['dmax'] = d
            if res['dmin'] is None or d < res['dmin']:
                res['dmin'] = d
            cip = code_cip(ligne['ARTCODEBARRE'])
            if not cip:
                res['sans_cip'] += 1
                continue
            cip = alias.get(cip, cip)
            exp = est_export(activites.get(code_tir(ligne['TIRCODE'])), ligne['TVATAUX'])
            e = res['cips'].setdefault((site, cip), [0, 0, None])
            e[1 if exp else 0] += q
            if q > 0 and not str(ligne['PCVNUM'] or '').startswith('AC'):
                if e[2] is None or d > e[2]:
                    e[2] = d
        wb.close()
    except Exception as ex:   # un fichier illisible est un REFUS d'écrire, pas un trou silencieux
        res['erreur'] = repr(ex)
    return res


def lister(stats):
    taches = []
    for dossier in ('', 'hors-perimetre'):
        base = os.path.join(stats, dossier) if dossier else stats
        if not os.path.isdir(base):
            continue
        for f in sorted(os.listdir(base)):
            m = RX_FICHIER.match(f)
            if m:
                taches.append((os.path.join(base, f), m.group(1), int(m.group(2))))
    return taches


def lire_activites(stats):
    """TIRCODE -> TIRACTIVITE, dans le plus récent clients_actifs_*.xlsx (1re occurrence gagne)."""
    fichiers = sorted(glob.glob(os.path.join(stats, 'total ventes', 'clients_actifs_*.xlsx')))
    if not fichiers:
        raise SystemExit('REFUS : aucun total ventes/clients_actifs_*.xlsx dans ' + stats)
    wb = openpyxl.load_workbook(fichiers[-1], read_only=True, data_only=True)
    it = wb.active.iter_rows(values_only=True)
    h = [str(x).strip() if x is not None else '' for x in next(it)]
    if 'TIRCODE' not in h or 'TIRACTIVITE' not in h:
        raise SystemExit('REFUS : colonnes TIRCODE / TIRACTIVITE absentes de ' + fichiers[-1])
    it_c, it_a = h.index('TIRCODE'), h.index('TIRACTIVITE')
    act = {}
    for r in it:
        if not r or r[it_c] is None:
            continue
        k = code_tir(r[it_c])
        if k not in act:
            act[k] = r[it_a]
    wb.close()
    return act, os.path.basename(fichiers[-1])


def lire_alias(stats):
    p = os.path.join(stats, 'pont-codes.json')
    return json.load(open(p, encoding='utf-8')).get('alias', {}) if os.path.exists(p) else {}


# ── Construction ────────────────────────────────────────────────────────────
def construire(stats, processus=4, journal=print):
    """-> (donnees, resume). Lève SystemExit('REFUS : ...') si un contrôle échoue."""
    activites, nom_clients = lire_activites(stats)
    alias = lire_alias(stats)
    taches = lister(stats)
    if not taches:
        raise SystemExit('REFUS : aucun fichier <CODE>_<MM>_2026.xlsx dans ' + stats)
    journal('%d fichiers à lire, clients : %s, %d codes dans le pont' % (len(taches), nom_clients, len(alias)))
    args = [(c, f, m, activites, alias) for c, f, m in taches]
    t0 = time.time()
    if processus > 1:
        with Pool(processus) as p:
            res = p.map(lire_fichier, args, chunksize=1)
    else:
        res = [lire_fichier(a) for a in args]
    journal('lecture : %.0f s' % (time.time() - t0))

    erreurs = [(r['fcode'], r['mois'], r['erreur']) for r in res if r['erreur']]
    if erreurs:
        raise SystemExit('REFUS : fichiers illisibles : %s' % erreurs[:5])

    mois_fichiers = {}
    for r in res:
        mois_fichiers[r['mois']] = mois_fichiers.get(r['mois'], 0) + 1
    dernier = max(mois_fichiers)
    manquants = [m for m in range(1, dernier + 1) if m not in mois_fichiers]
    if manquants:
        raise SystemExit('REFUS : aucun fichier pour le(s) mois %s (janvier à %d lus)' % (manquants, dernier))

    lignes_sm = {s: {m: 0 for m in range(1, dernier + 1)} for s in SITES}
    escale = sum(r['escale'] for r in res)
    inconnu, sans_cip = {}, 0
    for r in res:
        for s, n in r['sites'].items():
            lignes_sm[s][r['mois']] += n
        for s, n in r['inconnu'].items():
            inconnu[s] = inconnu.get(s, 0) + n
        sans_cip += r['sans_cip']
    depuis = {}
    for s in SITES:
        m = premier_mois_normal(lignes_sm[s])
        if m is None:
            raise SystemExit('REFUS : le site %s n\'a aucune ligne' % s)
        trous = [x for x in range(m, dernier + 1) if lignes_sm[s][x] == 0]
        if trous:
            raise SystemExit('REFUS : %s n\'a aucune ligne au(x) mois %s alors qu\'il vend depuis %d' % (s, trous, m))
        depuis[s] = m

    v = {}
    for r in res:
        for (s, cip), (uf, ue, dv) in r['cips'].items():
            if r['mois'] < depuis[s]:
                continue
            e = v.setdefault(cip, {}).setdefault(s, [0, 0, None])
            e[0] += uf
            e[1] += ue
            if dv and (e[2] is None or dv > e[2]):
                e[2] = dv
    sortie = {}
    for cip, par in v.items():
        for s, (uf, ue, dv) in par.items():
            uf, ue = int(round(uf)), int(round(ue))
            if uf == 0 and ue == 0:
                continue
            sortie.setdefault(cip, {})[s] = [uf, ue, dv or '']
    dmax = max(r['dmax'] for r in res if r['dmax'])
    annee = 2026
    donnees = {
        'de': '%d-%02d-01' % (annee, min(mois_fichiers)),
        'a': dmax,
        'genere': datetime.date.today().isoformat(),
        'regleExport': REGLE_EXPORT,
        'regleDepuis': REGLE_DEPUIS,
        'lignes': sum(r['lignes'] for r in res),
        'fichiers': len(res),
        'sites': {s: {'depuis': '%d-%02d' % (annee, depuis[s])} for s in SITES},
        'v': sortie,
    }
    resume = {'escale_lignes_a_part': escale, 'sites_inconnus': inconnu, 'lignes_sans_cip': sans_cip,
              'lignes_par_site_mois': lignes_sm, 'depuis': depuis, 'cip': len(sortie),
              'cases': sum(len(x) for x in sortie.values())}
    return donnees, resume


def controler_precedent(donnees, chemin):
    """Refuse si la fois précédente avait nettement plus de lignes, ou une facture plus récente."""
    if not os.path.exists(chemin):
        return
    try:
        txt = open(chemin, encoding='utf-8').read()
        ancien = json.loads(txt[txt.index('{'):txt.rindex('}') + 1])
    except Exception:
        return   # précédent illisible : rien à comparer (le signaler, pas bloquer)
    if ancien.get('lignes') and donnees['lignes'] < SEUIL_RELECTURE * ancien['lignes']:
        raise SystemExit('REFUS : %d lignes lues contre %d la fois précédente (moins de 80 %%)'
                         % (donnees['lignes'], ancien['lignes']))
    if ancien.get('a') and donnees['a'] < ancien['a']:
        raise SystemExit('REFUS : dernière facture lue %s, plus ancienne que la fois précédente (%s)'
                         % (donnees['a'], ancien['a']))


def ecrire(donnees, chemin):
    texte = ('// Ventes par établissement (France / export) depuis janvier — generate_etab_ventes.py\n'
             '// ⚠️ DONNÉE PROTÉGÉE : jamais dans git (.gitignore), déposée dans l\'espace Supabase fermé.\n'
             'window.ETAB_VENTES = ' + json.dumps(donnees, separators=(',', ':'), ensure_ascii=False) + ';\n')
    tmp = chemin + '.nouveau'
    with open(tmp, 'w', encoding='utf-8') as fh:
        fh.write(texte)
    os.replace(tmp, chemin)
    return len(texte.encode('utf-8'))


def main():
    donnees, resume = construire(STATS)
    controler_precedent(donnees, OUT)
    taille = ecrire(donnees, OUT)
    print('écrit %s : %d octets (%.2f Mo), %d produits, %d cases site×produit' % (OUT, taille, taille / 1e6, resume['cip'], resume['cases']))
    print('période %s -> %s ; depuis par site : %s' % (donnees['de'], donnees['a'], {s: donnees['sites'][s]['depuis'] for s in SITES}))
    print('à part (jamais dans le fichier) : %d lignes Escale ; lignes sans CIP : %d ; sources inconnues : %s'
          % (resume['escale_lignes_a_part'], resume['lignes_sans_cip'], resume['sites_inconnus'] or 'aucune'))
    if taille > 1.5e6:
        print('⚠️ fichier au-dessus de 1,5 Mo')


if __name__ == '__main__':
    try:
        main()
    except SystemExit as e:
        print(e)
        sys.exit(1)
